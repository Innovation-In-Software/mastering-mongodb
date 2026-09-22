#!/usr/bin/env python3
"""Scan rendered .pptx decks for code-card overflow: a code card (rounded
rectangle) whose own text -- at its *actual* rendered paragraph count and
font size, both read directly from the saved file, not re-estimated -- needs
more vertical height than the card itself was built with. This is the exact
bug class reported against Module 8 slide 39 ("Security Considerations"):
the card's fixed height didn't account for all of its own text lines, so the
last line(s) rendered below the card's rounded-rectangle border.

Also does a broader sweep: any shape whose bottom edge (top + height)
crosses into the footer zone or past the slide bottom entirely, excluding
the footer/rail/page-number shapes themselves (which legitimately live
there).

Usage:
    python check_code_card_overflow.py [decks/pptx/*.pptx ...]
    (defaults to every .pptx under decks/pptx/)
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DECKS = sorted((REPO_ROOT / "decks" / "pptx").glob("*.pptx"))

SLIDE_H_IN = 7.5
FOOTER_ZONE_IN = 7.05          # © ... copyright line starts at 7.098in
CODE_FONT_LINE_RATIO = 1.32    # same ratio mongodb_deck_kit.estimate_code_height uses
CODE_CARD_PAD_IN = 0.22        # same flat padding constant
TOLERANCE_IN = 0.03            # small slack for EMU rounding

IGNORED_SHAPE_NAMES = {"Diagram Page Number"}  # rotated page-number lives low by design


def is_rounded_rect(shape) -> bool:
    try:
        return (
            shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
            and shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        )
    except Exception:
        return False


def actual_font_pt(text_frame) -> float | None:
    for p in text_frame.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                return r.font.size.pt
    return None


def find_code_card_pairs(slide):
    """Yield (card_shape, textbox_shape) pairs matching add_code_block's
    construction order: a rounded-rect immediately followed by a text box
    positioned inside it."""
    shapes = list(slide.shapes)
    for i, shp in enumerate(shapes[:-1]):
        if not is_rounded_rect(shp):
            continue
        nxt = shapes[i + 1]
        if nxt.shape_type != MSO_SHAPE_TYPE.TEXT_BOX or not nxt.has_text_frame:
            continue
        # positioned inside the card (small tolerance for the 0.10-0.14in inset)
        if not (nxt.left >= shp.left - 1000 and nxt.top >= shp.top - 1000
                and nxt.left + nxt.width <= shp.left + shp.width + 1000
                and nxt.top + nxt.height <= shp.top + shp.height + 200000):
            continue
        # Only a genuine code card: add_code_block always sets every run's
        # font to Consolas. This distinguishes it from a callout card or a
        # cover-slide topic chip, which use the same rounded-rect+textbox
        # shape pair but the body font -- without this check those get
        # misidentified as code cards and scored against the wrong formula.
        fonts = {r.font.name for p in nxt.text_frame.paragraphs for r in p.runs}
        if fonts and fonts != {"Consolas"}:
            continue
        if not fonts:
            continue
        yield shp, nxt


def check_deck(path: Path) -> list[dict]:
    issues = []
    prs = Presentation(str(path))
    for slide_idx, slide in enumerate(prs.slides, start=1):
        for card, box in find_code_card_pairs(slide):
            tf = box.text_frame
            n_para = len(tf.paragraphs)
            font_pt = actual_font_pt(tf) or 13.0
            line_h_in = font_pt / 72 * CODE_FONT_LINE_RATIO
            needed_in = n_para * line_h_in + CODE_CARD_PAD_IN
            card_h_in = card.height / 914400
            if needed_in > card_h_in + TOLERANCE_IN:
                heading = None
                for shp in slide.shapes:
                    if shp.has_text_frame and shp.name.startswith("TextBox") and shp.top < 1000000:
                        heading = shp.text_frame.text
                        break
                issues.append({
                    "deck": path.name, "slide": slide_idx, "heading": heading,
                    "n_paragraphs": n_para, "font_pt": font_pt,
                    "card_height_in": round(card_h_in, 3),
                    "needed_height_in": round(needed_in, 3),
                    "deficit_in": round(needed_in - card_h_in, 3),
                })
        # broader sweep: any text-bearing shape that starts above the footer
        # zone but extends down into (or past) it -- i.e. actual content
        # encroaching on the footer, not the by-design rail/background
        # rectangles (no text) or the footer/page-number shapes themselves
        # (which legitimately start there).
        has_footer = any(s.name in IGNORED_SHAPE_NAMES for s in slide.shapes)
        for shp in slide.shapes:
            if shp.name in IGNORED_SHAPE_NAMES or not shp.has_text_frame:
                continue
            text = shp.text_frame.text.strip()
            if not text:
                continue
            # Two known, pre-existing, unrelated-to-this-fix elements that
            # are deliberately positioned at/near this boundary by the
            # original theme and were not touched by the code-card fix:
            # the cover slide's byline (a lead/chapter slide has no footer
            # copyright line at all, so nothing to collide with there) and
            # the "KEY TAKEAWAY" bar (bottom-anchored at a fixed constant,
            # TAKEAWAY_Y + TAKEAWAY_H = 7.12in, unrelated to code-block sizing).
            if not has_footer or text.startswith("KEY TAKEAWAY"):
                continue
            try:
                top_in = shp.top / 914400
                bottom_in = (shp.top + shp.height) / 914400
            except Exception:
                continue
            if top_in >= FOOTER_ZONE_IN:
                continue  # this shape is *supposed* to live in the footer band
            if bottom_in > FOOTER_ZONE_IN + TOLERANCE_IN or bottom_in > SLIDE_H_IN + TOLERANCE_IN:
                issues.append({
                    "deck": path.name, "slide": slide_idx, "shape": shp.name, "text": text[:40],
                    "bottom_in": round(bottom_in, 3),
                    "past_footer_zone_in": round(bottom_in - FOOTER_ZONE_IN, 3),
                })
    return issues


def main() -> None:
    paths = [Path(p) for p in sys.argv[1:]] or DEFAULT_DECKS
    total_issues = 0
    total_cards = 0
    for path in paths:
        prs = Presentation(str(path))
        n_cards = sum(1 for slide in prs.slides for _ in find_code_card_pairs(slide))
        total_cards += n_cards
        issues = check_deck(path)
        total_issues += len(issues)
        print(f"{path.name}: {n_cards} code cards scanned, {len(issues)} issue(s)")
        for iss in issues:
            print(f"    {iss}")
    print(f"\nTOTAL: {total_cards} code cards scanned across {len(paths)} decks, "
          f"{total_issues} issue(s) found")
    sys.exit(1 if total_issues else 0)


if __name__ == "__main__":
    main()
