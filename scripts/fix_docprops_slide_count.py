#!/usr/bin/env python3
"""Fix docProps/app.xml slide-count staleness in decks/pptx/*.pptx.

python-pptx does not update docProps/app.xml when slides are added via
prs.slides.add_slide() / prs.save() -- <Slides>, <Notes>, the HeadingPairs
"Slide Titles" count, and the TitlesOfParts vector (its size attribute and
its "Slide N" / "PowerPoint Presentation" placeholder entries) all go stale.
Real PowerPoint reads these on open and shows a "repair" / "found a problem
with content" dialog when they disagree with the actual slide count.
python-pptx, LibreOffice, and generic validators do not check this at all.

Ported unchanged (this logic is fully generic/content-agnostic) from the
sibling MD287 project's ``scripts/fix_docprops_slide_count.py``.

Usage:
    python fix_docprops_slide_count.py <pptx_path> [<pptx_path> ...]
    python fix_docprops_slide_count.py --check <pptx_path> [...]   # report only
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
import zipfile
from pathlib import Path

from pptx import Presentation

APP_XML_PATH = "docProps/app.xml"


def actual_slide_and_notes_count(pptx_path: Path) -> tuple[int, int]:
    prs = Presentation(str(pptx_path))
    total = len(prs.slides)
    with_notes = sum(
        1
        for slide in prs.slides
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
    )
    return total, with_notes


def read_app_xml(pptx_path: Path) -> str:
    with zipfile.ZipFile(pptx_path) as z:
        return z.read(APP_XML_PATH).decode("utf-8")


def build_fixed_app_xml(xml: str, actual_slides: int, actual_notes: int) -> tuple[str, dict]:
    report: dict = {}

    old_slides_match = re.search(r"<Slides>(\d+)</Slides>", xml)
    if not old_slides_match:
        raise ValueError("no <Slides> tag found")
    old_slides = int(old_slides_match.group(1))
    report["old_slides"] = old_slides
    report["new_slides"] = actual_slides

    xml = re.sub(r"<Slides>\d+</Slides>", f"<Slides>{actual_slides}</Slides>", xml)
    xml = re.sub(r"<Notes>\d+</Notes>", f"<Notes>{actual_notes}</Notes>", xml)
    xml = re.sub(
        r"(Slide Titles</vt:lpstr></vt:variant>\s*<vt:variant>\s*<vt:i4>)\d+(</vt:i4>)",
        r"\g<1>" + str(actual_slides) + r"\g<2>",
        xml,
    )

    m = re.search(
        r'(<TitlesOfParts>\s*<vt:vector size=")(\d+)("\s+baseType="lpstr">)(.*?)(</vt:vector>\s*</TitlesOfParts>)',
        xml,
        re.S,
    )
    if not m:
        raise ValueError("no <TitlesOfParts> vector found")
    old_size = int(m.group(2))
    inner = m.group(4)

    entries = re.findall(r"<vt:lpstr>.*?</vt:lpstr>", inner)
    if len(entries) != old_size:
        raise ValueError(f"TitlesOfParts entry count ({len(entries)}) != size attr ({old_size})")

    # Non-slide entries (fonts + theme, always listed first) are however many
    # entries are left after the old declared slide count -- NOT whichever
    # entries happen to match a "Slide N" text pattern.
    non_slide_count = old_size - old_slides
    if non_slide_count < 0:
        raise ValueError(f"old <Slides>{old_slides}</Slides> exceeds TitlesOfParts size {old_size}")
    non_slide_entries = entries[:non_slide_count]
    old_slide_entries = entries[non_slide_count:]
    report["non_slide_titles"] = non_slide_count
    report["old_slide_titles"] = len(old_slide_entries)
    report["new_slide_titles"] = actual_slides

    placeholder = (
        "<vt:lpstr>Slide {}</vt:lpstr>"
        if old_slide_entries and re.match(r"<vt:lpstr>Slide \d+</vt:lpstr>", old_slide_entries[-1])
        else None
    )
    new_size = non_slide_count + actual_slides
    if placeholder:
        new_slide_entries = "".join(placeholder.format(i) for i in range(1, actual_slides + 1))
    else:
        new_slide_entries = "<vt:lpstr>PowerPoint Presentation</vt:lpstr>" * actual_slides
    new_inner = "".join(non_slide_entries) + new_slide_entries

    xml = (
        xml[: m.start()]
        + m.group(1)
        + str(new_size)
        + m.group(3)
        + new_inner
        + m.group(5)
        + xml[m.end() :]
    )

    return xml, report


def rewrite_pptx_with_new_app_xml(pptx_path: Path, new_app_xml: str) -> None:
    temporary = pptx_path.with_suffix(".docprops-fix.pptx")
    temporary.unlink(missing_ok=True)
    with zipfile.ZipFile(pptx_path) as src, zipfile.ZipFile(
        temporary, "w", zipfile.ZIP_DEFLATED
    ) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == APP_XML_PATH:
                data = new_app_xml.encode("utf-8")
            dst.writestr(item, data)
    # Reopen before replacing the original to catch package/XML errors.
    Presentation(str(temporary))
    shutil.move(str(temporary), str(pptx_path))


def process(pptx_path: Path, *, check_only: bool) -> None:
    actual_slides, actual_notes = actual_slide_and_notes_count(pptx_path)
    xml = read_app_xml(pptx_path)
    old_slides_match = re.search(r"<Slides>(\d+)</Slides>", xml)
    old_slides = int(old_slides_match.group(1)) if old_slides_match else None

    if old_slides == actual_slides:
        print(f"OK       {pptx_path.name}: docProps already matches ({actual_slides} slides)")
        return

    if check_only:
        print(
            f"MISMATCH {pptx_path.name}: docProps says {old_slides}, actual is {actual_slides}"
        )
        return

    new_xml, report = build_fixed_app_xml(xml, actual_slides, actual_notes)
    rewrite_pptx_with_new_app_xml(pptx_path, new_xml)
    print(
        f"FIXED    {pptx_path.name}: slides {report['old_slides']}->{report['new_slides']}, "
        f"notes ->{actual_notes}, slide-titles {report['old_slide_titles']}->{report['new_slide_titles']}"
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx_paths", nargs="+", type=Path)
    parser.add_argument("--check", action="store_true", help="report mismatches, don't fix")
    args = parser.parse_args()

    had_error = False
    for path in args.pptx_paths:
        try:
            process(path, check_only=args.check)
        except Exception as exc:  # noqa: BLE001
            had_error = True
            print(f"ERROR    {path.name}: {exc}", file=sys.stderr)

    if had_error:
        sys.exit(1)


if __name__ == "__main__":
    main()
