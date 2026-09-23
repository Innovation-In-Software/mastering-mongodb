#!/usr/bin/env python3
"""
Kahoot bulk spreadsheet uploader (Playwright).

Kahoot has no public bulk-create API. This script opens Chromium, waits for
you to log in once, then for each .xlsx:

  Create → Import spreadsheet → upload → set title → save

Default source is the course `kahoot/` folder (Modules 1–8).

Usage examples (from repo root):

  python scripts/kahoot_bulk_upload/kahoot_bulk_upload.py --max 2
  python scripts/kahoot_bulk_upload/kahoot_bulk_upload.py
  python scripts/kahoot_bulk_upload/kahoot_bulk_upload.py --source kahoot/Kahoot_Module_1.xlsx

Important:
- Automates the visible browser UI only (no login/CAPTCHA bypass).
- Spreadsheet import may require a Kahoot plan that includes that feature.
- Kahoot UI changes can break selectors; failures write screenshot+HTML under logs/.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from playwright.sync_api import (
    BrowserContext,
    Locator,
    Page,
    Playwright,
    TimeoutError as PlaywrightTimeoutError,
    sync_playwright,
)

TOOL_DIR = Path(__file__).resolve().parent
REPO = TOOL_DIR.parents[1]
DEFAULT_SOURCE = REPO / "kahoot"
# Reuse the instructor Chrome profile that is already logged in to Kahoot.
_SHARED_PROFILE = Path(
    r"D:\Current_work\Innovation in Software\MD287 Spring Boot Microservices with OpenShift AI and DevOps_Bank of America\scripts\kahoot_bulk_upload\browser_profile"
)
PROFILE_DIR = _SHARED_PROFILE if (_SHARED_PROFILE / "Local State").exists() else (TOOL_DIR / "browser_profile")
LOG_DIR = TOOL_DIR / "logs"
RESULTS_FILE = TOOL_DIR / "upload_results.csv"
PREPARED_DIR = TOOL_DIR / "prepared"
TRACKER_PATH = REPO / "kahoot" / "Kahoot_Quiz_Share_Links.xlsx"

CREATE_URL = "https://create.kahoot.it/creator"
STEP_PAUSE_MS = 1000
ACTION_TIMEOUT_MS = 20_000

# Official Kahoot spreadsheet import headers (as of Kahoot help docs).
KAHOOT_HEADERS = [
    "Question",
    "Answer 1",
    "Answer 2",
    "Answer 3",
    "Answer 4",
    "Time limit (seconds)",
    "Correct answer(s)",
]
QUESTION_MAX = 95
ANSWER_MAX = 60
ALLOWED_TIMES = {5, 10, 20, 30, 60, 120}

if str(TOOL_DIR.parent) not in sys.path:
    sys.path.insert(0, str(TOOL_DIR.parent))
from mongodb_kahoot_questions import MODULE_TITLES  # noqa: E402


def log(message: str) -> None:
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {message}", flush=True)


def quiz_title(path: Path, prefix: str) -> str:
    match = re.search(r"Module[_\s-]?(\d+)", path.stem, re.I)
    if match:
        num = int(match.group(1))
        name = MODULE_TITLES.get(num, path.stem)
        title = f"{prefix}Module {num} — {name}".strip()
    else:
        title = f"{prefix}{path.stem.replace('_', ' ')}".strip()
    return re.sub(r"\s+", " ", title)[:95]


def discover_xlsx(source: Path, include_day2: bool) -> list[Path]:
    if source.is_file() and source.suffix.lower() == ".xlsx":
        return [source]
    if not source.exists():
        raise FileNotFoundError(f"Source not found: {source}")
    files = sorted(p for p in source.rglob("*.xlsx") if p.is_file())
    # Only module quiz banks — never the share-links tracker or other workbooks.
    module_files = [p for p in files if re.match(r"(?i)^Kahoot_Module_\d+\.xlsx$", p.name)]
    if include_day2:
        module_files.extend(
            p for p in files if p.name.lower() == "kahoot_day2.xlsx" and p not in module_files
        )
    module_files = [p for p in module_files if not p.name.startswith("~$")]
    return sorted(module_files, key=lambda p: (module_from_stem(p.stem), p.as_posix().lower()))


def module_from_stem(stem: str) -> int:
    match = re.search(r"(\d+)", stem)
    return int(match.group(1)) if match else 0


def prepare_for_kahoot(src: Path, dest: Path) -> list[str]:
    """Rewrite workbook to Kahoot's template headers; truncate over-limit text."""
    warnings: list[str] = []
    wb_in = load_workbook(src, data_only=True)
    ws_in = wb_in.active

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "Sheet1"
    ws_out.append(KAHOOT_HEADERS)

    for row_idx, row in enumerate(ws_in.iter_rows(min_row=2, values_only=True), start=2):
        if not row or row[0] is None:
            continue
        cells = list(row) + [None] * (7 - len(row))
        question = str(cells[0] or "").strip()
        answers = [str(cells[i] or "").strip() for i in range(1, 5)]
        time_val = cells[5]
        correct = cells[6]

        if len(question) > QUESTION_MAX:
            warnings.append(
                f"{src.name} row {row_idx}: question truncated {len(question)}→{QUESTION_MAX}"
            )
            question = question[:QUESTION_MAX]
        for i, ans in enumerate(answers):
            if len(ans) > ANSWER_MAX:
                warnings.append(
                    f"{src.name} row {row_idx}: answer {i+1} truncated {len(ans)}→{ANSWER_MAX}"
                )
                answers[i] = ans[:ANSWER_MAX]

        try:
            time_int = int(time_val)
        except (TypeError, ValueError):
            time_int = 30
            warnings.append(f"{src.name} row {row_idx}: invalid time -> 30")
        if time_int not in ALLOWED_TIMES:
            warnings.append(f"{src.name} row {row_idx}: time {time_int} -> 30")
            time_int = 30

        if correct is None or str(correct).strip() == "":
            raise ValueError(f"{src.name} row {row_idx}: missing Correct value")
        correct_str = str(correct).strip()

        ws_out.append([question, *answers, time_int, correct_str])

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(dest)
    return warnings


def click_first(page: Page, candidates: Iterable[Locator], description: str) -> None:
    last_error: Exception | None = None
    for locator in candidates:
        try:
            candidate = locator.first
            candidate.wait_for(state="visible", timeout=4_000)
            try:
                candidate.click(timeout=ACTION_TIMEOUT_MS)
            except Exception:
                candidate.click(timeout=ACTION_TIMEOUT_MS, force=True)
            page.wait_for_timeout(STEP_PAUSE_MS)
            log(f"Clicked: {description}")
            return
        except Exception as exc:
            last_error = exc
    raise RuntimeError(f"Could not find/click {description}. Last error: {last_error}")


def fill_first(page: Page, candidates: Iterable[Locator], value: str, description: str) -> None:
    last_error: Exception | None = None
    for locator in candidates:
        try:
            candidate = locator.first
            candidate.wait_for(state="visible", timeout=4_000)
            candidate.click(timeout=ACTION_TIMEOUT_MS)
            page.wait_for_timeout(200)
            # Clear any AI-prefilled draft, then type the bootcamp title.
            try:
                candidate.fill("")
            except Exception:
                page.keyboard.press("Control+A")
                page.keyboard.press("Backspace")
            candidate.fill(value, timeout=ACTION_TIMEOUT_MS)
            page.wait_for_timeout(300)
            log(f"Filled: {description}")
            return
        except Exception as exc:
            last_error = exc
    raise RuntimeError(f"Could not fill {description}. Last error: {last_error}")


def dismiss_blocking_overlays(page: Page) -> None:
    """Try to clear cookie/consent overlays that gray-out the login form."""
    try:
        page.keyboard.press("Escape")
        page.wait_for_timeout(400)
    except Exception:
        pass

    patterns = [
        r"accept all",
        r"accept cookies",
        r"allow all",
        r"i agree",
        r"^agree$",
        r"got it",
        r"^ok$",
        r"continue",
        r"^close$",
        r"maybe later",
        r"not now",
        r"^skip$",
        r"no thanks",
    ]
    for pattern in patterns:
        try:
            btn = page.get_by_role("button", name=re.compile(pattern, re.I)).first
            if btn.is_visible(timeout=800):
                btn.click(timeout=2_000)
                page.wait_for_timeout(500)
                log(f"Dismissed overlay via: {pattern}")
                return
        except Exception:
            continue

    # OneTrust / common cookie managers
    for selector in [
        "#onetrust-accept-btn-handler",
        "button#accept-recommended-btn-handler",
        '[data-testid*="cookie"] button',
        'button:has-text("Accept")',
    ]:
        try:
            loc = page.locator(selector).first
            if loc.is_visible(timeout=500):
                loc.click(timeout=2_000)
                page.wait_for_timeout(500)
                log(f"Dismissed overlay via selector: {selector}")
                return
        except Exception:
            continue


def _looks_logged_out(page: Page) -> bool:
    url = page.url.lower()
    if "/auth/login" in url or "/login" in url or "signin" in url:
        return True
    try:
        if page.locator('input[type="password"]').count() > 0:
            return True
    except Exception:
        pass
    return False


def wait_for_manual_login(page: Page, timeout_sec: int = 1800) -> None:
    page.goto(CREATE_URL, wait_until="domcontentloaded", timeout=90_000)
    page.wait_for_timeout(2_000)
    dismiss_blocking_overlays(page)

    print("\n" + "=" * 72, flush=True)
    print("A Chrome window should be open on Kahoot.", flush=True)
    print("If the page looks grayed out:", flush=True)
    print("  1) Press Escape", flush=True)
    print("  2) Click Accept / Accept all on any cookie banner", flush=True)
    print("  3) Then log in (Google/Microsoft/email).", flush=True)
    print("This script waits quietly and will NOT reload the page while you type.", flush=True)
    print(f"Waiting up to {timeout_sec // 60} minutes...", flush=True)
    print("=" * 72 + "\n", flush=True)

    deadline = datetime.now().timestamp() + timeout_sec
    last_dismiss = 0.0
    while datetime.now().timestamp() < deadline:
        # Periodically retry dismissing overlays, but never navigate away.
        now = datetime.now().timestamp()
        if now - last_dismiss > 15:
            dismiss_blocking_overlays(page)
            last_dismiss = now

        url = page.url.lower()
        if "/auth/login" not in url and "login" not in url and not _looks_logged_out(page):
            log("Login detected — starting bulk import.")
            return
        if "create.kahoot.it" in url and "/auth/" not in url and "login" not in url:
            # Creator/home without auth path.
            log(f"Login detected (URL: {page.url}) — starting bulk import.")
            return
        page.wait_for_timeout(2_000)

    raise TimeoutError(
        "Timed out waiting for Kahoot login. Log in in the browser window and run again."
    )


def open_blank_creator(page: Page) -> None:
    page.goto(CREATE_URL, wait_until="domcontentloaded", timeout=90_000)
    page.wait_for_timeout(2_000)

    try:
        kahoot_option = page.get_by_text(re.compile(r"^kahoot!?$", re.I)).first
        if kahoot_option.is_visible(timeout=2_000):
            kahoot_option.click()
            page.wait_for_timeout(STEP_PAUSE_MS)
    except Exception:
        pass

    # Blank / start from scratch if a template chooser appears.
    for pattern in [r"blank", r"start from scratch", r"create.*blank"]:
        try:
            btn = page.get_by_role("button", name=re.compile(pattern, re.I)).first
            if btn.is_visible(timeout=1_500):
                btn.click()
                page.wait_for_timeout(STEP_PAUSE_MS)
                break
        except Exception:
            continue

    if "creator" not in page.url.lower():
        try:
            click_first(
                page,
                [
                    page.get_by_role("button", name=re.compile(r"^create", re.I)),
                    page.get_by_role("link", name=re.compile(r"^create", re.I)),
                    page.get_by_text(re.compile(r"^create$", re.I)),
                ],
                "Create",
            )
        except Exception:
            pass

    page.wait_for_timeout(1_500)


def _add_question_locators(page: Page) -> list[Locator]:
    return [
        page.locator('[data-functional-selector="add-question-button"]'),
        page.get_by_role("button", name=re.compile(r"add question", re.I)),
        page.get_by_text(re.compile(r"add question", re.I)),
        page.locator('[data-functional-selector*="add-question"]'),
        page.locator('button:has-text("Add question")'),
        page.get_by_role("button", name=re.compile(r"^\+?\s*add$", re.I)),
        page.locator('button:has-text("+ Add")'),
    ]


def _import_locators(page: Page) -> list[Locator]:
    return [
        page.get_by_role("menuitem", name=re.compile(r"^import", re.I)),
        page.get_by_role("button", name=re.compile(r"^import", re.I)),
        page.get_by_text(re.compile(r"^import$", re.I)),
        page.get_by_text(re.compile(r"import questions", re.I)),
        page.locator('button:has-text("Import")'),
        page.locator('[role="menuitem"]:has-text("Import")'),
    ]


def open_spreadsheet_import(page: Page) -> None:
    # The Add menu sometimes does not open on the first click (creator still settling).
    last_error: Exception | None = None
    for attempt in range(3):
        try:
            blocker = page.locator(
                '[data-functional-selector="add-question-interaction-blocker"]'
            )
            if blocker.count() > 0 and blocker.first.is_visible(timeout=500):
                blocker.first.wait_for(state="hidden", timeout=8_000)
        except Exception:
            pass
        click_first(page, _add_question_locators(page), "Add question")
        try:
            click_first(page, _import_locators(page), "Import")
            last_error = None
            break
        except Exception as exc:
            last_error = exc
            log(f"Import menu not open (attempt {attempt + 1}); retrying Add")
            try:
                page.keyboard.press("Escape")
            except Exception:
                pass
            page.wait_for_timeout(800)
    if last_error is not None:
        raise last_error

    try:
        spreadsheet = page.get_by_text(
            re.compile(r"import.*spreadsheet|spreadsheet.*import", re.I)
        ).first
        if spreadsheet.is_visible(timeout=2_500):
            spreadsheet.click()
            page.wait_for_timeout(STEP_PAUSE_MS)
            log("Clicked: Import spreadsheet")
    except Exception:
        pass


def _question_title_text(page: Page) -> str:
    """Return the currently selected question title text (may be empty)."""
    try:
        title = page.locator('[data-functional-selector="question-title__input"]').first
        if title.count() == 0:
            return ""
        # contenteditable often exposes text via inner_text
        text = (title.inner_text(timeout=2_000) or "").strip()
        # Placeholder-only nodes can still report empty / whitespace.
        return text
    except Exception:
        return ""


def _selected_question_is_empty(page: Page) -> bool:
    title = _question_title_text(page)
    if title:
        return False
    # Require answers to also be blank so we never delete a real imported question
    # that is still hydrating its title.
    try:
        answers = page.locator('[data-functional-selector="question-answer__input"]')
        count = min(answers.count(), 4)
        if count == 0:
            return True
        for i in range(count):
            text = (answers.nth(i).inner_text(timeout=1_500) or "").strip()
            if text:
                return False
        return True
    except Exception:
        return title == ""


def _confirm_delete_if_needed(page: Page) -> None:
    for pattern in [r"^delete$", r"^remove$", r"^confirm$", r"^yes$"]:
        try:
            button = page.get_by_role("button", name=re.compile(pattern, re.I)).first
            if button.is_visible(timeout=1_200):
                button.click(timeout=ACTION_TIMEOUT_MS)
                page.wait_for_timeout(800)
                log(f"Confirmed delete via: {pattern}")
                return
        except Exception:
            continue


def remove_empty_questions(page: Page, max_passes: int = 5) -> int:
    """Delete blank/Untitled starter questions left before or after spreadsheet import.

    Kahoot often keeps an empty Question 1 ("Start typing your question") so a
    10-row spreadsheet becomes 11 cards with the first one blank.
    """
    removed = 0
    page.wait_for_timeout(1_500)

    for _ in range(max_passes):
        blocks = page.locator('[data-functional-selector^="sidebar-block__kahoot-block-"]')
        count = blocks.count()
        if count == 0:
            break

        deleted_this_pass = False
        # Walk from the start so the blank starter (usually block 0) is removed first.
        for index in range(count):
            try:
                block = blocks.nth(index)
                block.click(timeout=ACTION_TIMEOUT_MS)
                page.wait_for_timeout(700)
            except Exception:
                continue

            if not _selected_question_is_empty(page):
                continue

            # Prefer sidebar delete on the selected card, then right-panel delete.
            deleted = False
            for selector in [
                '[data-functional-selector="sidebar__remove"]',
                '[data-functional-selector="right-panel-delete-button"]',
                'button[aria-label="Delete"]',
            ]:
                try:
                    btn = page.locator(selector).first
                    if btn.is_visible(timeout=1_500):
                        btn.click(timeout=ACTION_TIMEOUT_MS)
                        page.wait_for_timeout(500)
                        _confirm_delete_if_needed(page)
                        removed += 1
                        deleted = True
                        deleted_this_pass = True
                        log(f"Removed empty question at sidebar index {index}")
                        page.wait_for_timeout(800)
                        break
                except Exception:
                    continue

            if deleted:
                # Restart scan because indexes shift after deletion.
                break

        if not deleted_this_pass:
            break

    if removed:
        log(f"Removed {removed} empty question(s) after import")
    else:
        log("No empty questions found to remove")

    # Ensure Question 1 is a real imported question.
    blocks = page.locator('[data-functional-selector^="sidebar-block__kahoot-block-"]')
    if blocks.count() > 0:
        blocks.first.click(timeout=ACTION_TIMEOUT_MS)
        page.wait_for_timeout(800)
        if _selected_question_is_empty(page):
            raise RuntimeError(
                "Question 1 is still empty after import cleanup. "
                "Check the spreadsheet and Kahoot import result."
            )
        log(f"Question 1 OK: {_question_title_text(page)[:80]}")
    return removed


def upload_spreadsheet(page: Page, file_path: Path) -> None:
    inputs = page.locator('input[type="file"]')
    if inputs.count() > 0:
        inputs.last.set_input_files(str(file_path.resolve()))
        log(f"Selected file: {file_path.name}")
    else:
        with page.expect_file_chooser(timeout=ACTION_TIMEOUT_MS) as chooser_info:
            click_first(
                page,
                [
                    page.get_by_role(
                        "button", name=re.compile(r"select file|choose file|browse", re.I)
                    ),
                    page.get_by_text(re.compile(r"select file|choose file|browse", re.I)),
                ],
                "Select file",
            )
        chooser_info.value.set_files(str(file_path.resolve()))
        log(f"Selected file: {file_path.name}")

    page.wait_for_timeout(1_000)

    click_first(
        page,
        [
            page.get_by_role("button", name=re.compile(r"^upload$", re.I)),
            page.get_by_text(re.compile(r"^upload$", re.I)),
            page.locator('button:has-text("Upload")'),
        ],
        "Upload",
    )

    page.wait_for_timeout(4_000)

    for pattern in [r"continue", r"add questions", r"import questions", r"done"]:
        try:
            button = page.get_by_role("button", name=re.compile(pattern, re.I)).first
            if button.is_visible(timeout=2_000):
                button.click()
                page.wait_for_timeout(2_000)
                log(f"Clicked post-upload: {pattern}")
                break
        except Exception:
            continue


def set_title_and_save(page: Page, title: str) -> str:
    """Set title in Kahoot settings (current UI) and save the kahoot."""
    # Prefer the top-bar title button ("Enter kahoot title…") / settings entry.
    opened_settings = False
    for locator in [
        page.get_by_role(
            "button", name=re.compile(r"enter kahoot title|untitled|kahoot title", re.I)
        ),
        page.locator('button[class*="TitleButton"]'),
        page.get_by_role("button", name=re.compile(r"settings", re.I)),
        page.locator('[data-functional-selector*="settings"]'),
        page.get_by_text(re.compile(r"^settings$", re.I)),
    ]:
        try:
            candidate = locator.first
            if candidate.is_visible(timeout=2_000):
                candidate.click(timeout=ACTION_TIMEOUT_MS)
                page.wait_for_timeout(STEP_PAUSE_MS)
                opened_settings = True
                log("Opened: Kahoot settings / title")
                break
        except Exception:
            continue

    if not opened_settings:
        raise RuntimeError("Could not open Kahoot settings/title dialog")

    # Settings dialog: label "Title" with a textbox (AI may already fill a draft).
    title_candidates = [
        page.get_by_label(re.compile(r"^title$", re.I)),
        page.get_by_role("textbox", name=re.compile(r"^title$", re.I)),
        page.locator('label:has-text("Title")').locator("xpath=following::input[1]"),
        page.locator('label:has-text("Title")').locator("xpath=following::textarea[1]"),
        page.locator('[data-functional-selector*="title"] input'),
        page.locator('[data-functional-selector*="title"] textarea'),
        page.locator('input[name*="title" i]'),
        page.locator('textarea[name*="title" i]'),
    ]
    fill_first(page, title_candidates, title, "quiz title")

    # Prefer Private for instructor banks when the control is present.
    try:
        private = page.get_by_role("radio", name=re.compile(r"private", re.I)).first
        if private.is_visible(timeout=1_500):
            private.check(timeout=ACTION_TIMEOUT_MS)
            log("Set visibility: Private")
    except Exception:
        try:
            page.get_by_text(re.compile(r"^private$", re.I)).first.click(timeout=2_000)
            log("Clicked: Private")
        except Exception:
            pass

    # Close settings with Done.
    click_first(
        page,
        [
            page.get_by_role("button", name=re.compile(r"^done$", re.I)),
            page.get_by_text(re.compile(r"^done$", re.I)),
            page.locator('button:has-text("Done")'),
        ],
        "Done (settings)",
    )
    page.wait_for_timeout(1_000)

    # Capture creator UUID from the editor URL before Save redirects to My Library.
    pre_save_url = page.url.strip()

    # Save the kahoot from the top bar.
    click_first(
        page,
        [
            page.get_by_role("button", name=re.compile(r"^save$", re.I)),
            page.get_by_text(re.compile(r"^save$", re.I)),
            page.locator('[data-functional-selector*="save"]'),
            page.locator('button:has-text("Save")'),
        ],
        "Save",
    )
    page.wait_for_timeout(2_000)

    # Confirm any follow-up save dialogs.
    for _ in range(3):
        clicked = False
        for pattern in [r"^save$", r"^done$", r"^continue$", r"^confirm$"]:
            try:
                button = page.get_by_role("button", name=re.compile(pattern, re.I)).first
                if button.is_visible(timeout=1_500):
                    button.click()
                    page.wait_for_timeout(1_500)
                    clicked = True
                    break
            except Exception:
                continue
        if not clicked:
            break

    page.wait_for_timeout(1_500)
    return pre_save_url


_CREATOR_UUID_RE = re.compile(
    r"https?://create\.kahoot\.it/creator/([0-9a-fA-F-]{36})",
    re.I,
)
_DETAILS_UUID_RE = re.compile(
    r"https?://create\.kahoot\.it/(?:v2/)?details/([0-9a-fA-F-]{36})",
    re.I,
)


def _is_per_quiz_url(url: str) -> bool:
    """True for a specific kahoot creator/details URL — not the library folder."""
    if not url:
        return False
    lower = url.lower()
    if "/my-library/" in lower:
        return False
    return bool(_CREATOR_UUID_RE.search(url) or _DETAILS_UUID_RE.search(url))


def _normalize_creator_url(url: str) -> str:
    """Prefer create.kahoot.it/creator/{uuid} when we can extract a UUID."""
    if not url:
        return ""
    match = _CREATOR_UUID_RE.search(url) or _DETAILS_UUID_RE.search(url)
    if match:
        return f"https://create.kahoot.it/creator/{match.group(1)}"
    return url.strip() if _is_per_quiz_url(url) else ""


def capture_quiz_urls(page: Page, pre_save_url: str = "") -> tuple[str, str]:
    """Return (creator_url, share_url) after a successful save.

    Kahoot often redirects to My Library after Save, so we prefer a UUID captured
    from the creator URL *before* save, then wait briefly for /creator/{uuid}.
    """
    creator_url = _normalize_creator_url(pre_save_url) or _normalize_creator_url(page.url)

    # After save, wait briefly for a stable per-quiz creator URL before library redirect.
    if not creator_url:
        deadline = datetime.now().timestamp() + 8
        while datetime.now().timestamp() < deadline:
            candidate = _normalize_creator_url(page.url)
            if candidate:
                creator_url = candidate
                break
            page.wait_for_timeout(400)

    share_url = ""

    # Prefer an explicit share/copy-link control when Kahoot shows one.
    for pattern in [
        r"copy link",
        r"share link",
        r"invite",
        r"^share$",
    ]:
        try:
            btn = page.get_by_role("button", name=re.compile(pattern, re.I)).first
            if not btn.is_visible(timeout=1_200):
                continue
            btn.click(timeout=ACTION_TIMEOUT_MS)
            page.wait_for_timeout(800)
            for locator in [
                page.locator('input[type="text"]'),
                page.locator("input"),
                page.locator('a[href*="kahoot"]'),
            ]:
                try:
                    count = min(locator.count(), 8)
                    for i in range(count):
                        el = locator.nth(i)
                        href = el.get_attribute("href")
                        value = el.input_value() if el.evaluate("e => 'value' in e") else None
                        candidate = (href or value or "").strip()
                        if candidate.startswith("http") and "kahoot" in candidate.lower():
                            if _is_per_quiz_url(candidate) or "play.kahoot.it" in candidate.lower():
                                share_url = candidate
                                break
                    if share_url:
                        break
                except Exception:
                    continue
            if share_url:
                break
        except Exception:
            continue

    if not share_url and creator_url:
        share_url = creator_url
    return creator_url, share_url


def capture_failure(page: Page, file_path: Path) -> tuple[str, str]:
    LOG_DIR.mkdir(exist_ok=True)
    safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", file_path.stem)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    screenshot = LOG_DIR / f"{safe_name}_{stamp}.png"
    html = LOG_DIR / f"{safe_name}_{stamp}.html"
    try:
        page.screenshot(path=str(screenshot), full_page=True)
    except Exception:
        screenshot = Path("")
    try:
        html.write_text(page.content(), encoding="utf-8")
    except Exception:
        html = Path("")
    return str(screenshot), str(html)


def append_result(filename: str, title: str, status: str, details: str = "") -> None:
    new_file = not RESULTS_FILE.exists()
    with RESULTS_FILE.open("a", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        if new_file:
            writer.writerow(["timestamp", "filename", "title", "status", "details"])
        writer.writerow(
            [datetime.now().isoformat(timespec="seconds"), filename, title, status, details]
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Bulk-create Kahoot quizzes from Mastering MongoDB Excel banks via Playwright."
    )
    parser.add_argument(
        "--source",
        type=Path,
        default=DEFAULT_SOURCE,
        help=f"Folder or file with .xlsx quizzes (default: {DEFAULT_SOURCE})",
    )
    parser.add_argument(
        "--max",
        type=int,
        default=None,
        help="Process only the first N files (use 2 for a smoke test).",
    )
    parser.add_argument(
        "--title-prefix",
        default="Mastering MongoDB — ",
        help="Prefix for Kahoot titles (default: 'Mastering MongoDB — ').",
    )
    parser.add_argument(
        "--include-day2",
        action="store_true",
        help="Also upload Kahoot_Day2.xlsx if present.",
    )
    parser.add_argument(
        "--prepare-only",
        action="store_true",
        help="Only rewrite spreadsheets to Kahoot template headers; do not open a browser.",
    )
    parser.add_argument(
        "--headless",
        action="store_true",
        help="Run Chromium headless (not recommended for first login).",
    )
    parser.add_argument(
        "--skip-uploaded",
        action="store_true",
        default=True,
        help="Skip modules already marked Uploaded in Kahoot_Quiz_Share_Links.xlsx (default: on).",
    )
    parser.add_argument(
        "--reupload-all",
        action="store_true",
        help="Upload every module even if already marked Uploaded (creates duplicates).",
    )
    return parser.parse_args()


def _uploaded_modules_from_tracker() -> set[int]:
    if not TRACKER_PATH.exists():
        return set()
    try:
        wb = load_workbook(TRACKER_PATH, read_only=True)
        ws = wb["Quiz Share Links"] if "Quiz Share Links" in wb.sheetnames else wb.active
        uploaded: set[int] = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            try:
                module = int(row[0])
            except (TypeError, ValueError):
                continue
            status = str(row[7] or "").strip().lower() if len(row) > 7 else ""
            if status == "uploaded":
                uploaded.add(module)
        return uploaded
    except Exception:
        return set()


def run(playwright: Playwright | None, args: argparse.Namespace) -> int:
    source = args.source if args.source.is_absolute() else (REPO / args.source)
    files = discover_xlsx(source, include_day2=args.include_day2)
    if args.max is not None:
        files = files[: args.max]

    if not files:
        print(f"No .xlsx files found under: {source}")
        return 1

    PREPARED_DIR.mkdir(parents=True, exist_ok=True)
    LOG_DIR.mkdir(exist_ok=True)

    prepared: list[tuple[Path, Path, str]] = []
    for src in files:
        title = quiz_title(src, args.title_prefix)
        dest = PREPARED_DIR / src.name
        warnings = prepare_for_kahoot(src, dest)
        for w in warnings:
            log(f"WARN: {w}")
        prepared.append((src, dest, title))

    skip_uploaded = args.skip_uploaded and not args.reupload_all
    if skip_uploaded:
        from quiz_url_tracker import module_from_filename

        already = _uploaded_modules_from_tracker()
        before = len(prepared)
        prepared = [
            item
            for item in prepared
            if module_from_filename(item[0].name) not in already
        ]
        skipped = before - len(prepared)
        if skipped:
            log(f"Skipping {skipped} already-uploaded module(s): {sorted(already)}")

    log(f"Prepared {len(prepared)} spreadsheet(s) under {PREPARED_DIR}")
    if args.prepare_only:
        return 0
    if not prepared:
        log("Nothing left to upload.")
        return 0

    assert playwright is not None
    PROFILE_DIR.mkdir(exist_ok=True)
    # Prefer installed Google Chrome (more reliable for Kahoot login overlays).
    # Falls back to bundled Chromium if Chrome is not installed.
    launch_kwargs = dict(
        user_data_dir=str(PROFILE_DIR),
        headless=args.headless,
        viewport={"width": 1400, "height": 900},
        accept_downloads=True,
        args=[
            "--disable-blink-features=AutomationControlled",
            "--disable-dev-shm-usage",
        ],
        ignore_default_args=["--enable-automation"],
    )
    try:
        context = playwright.chromium.launch_persistent_context(
            channel="chrome",
            **launch_kwargs,
        )
        log("Launched Google Chrome for Kahoot.")
    except Exception as exc:
        log(f"Chrome channel unavailable ({exc}); falling back to Chromium.")
        context = playwright.chromium.launch_persistent_context(**launch_kwargs)

    context.set_default_timeout(ACTION_TIMEOUT_MS)
    page = context.pages[0] if context.pages else context.new_page()

    wait_for_manual_login(page)

    successes = 0
    for index, (src, prepared_path, title) in enumerate(prepared, start=1):
        log(f"[{index}/{len(prepared)}] Starting: {src.name} -> {title}")
        try:
            open_blank_creator(page)
            open_spreadsheet_import(page)
            upload_spreadsheet(page, prepared_path)
            remove_empty_questions(page)
            pre_save_url = set_title_and_save(page, title)
            creator_url, share_url = capture_quiz_urls(page, pre_save_url=pre_save_url)
            append_result(
                src.name,
                title,
                "SUCCESS",
                f"creator_url={creator_url}; share_url={share_url}",
            )
            try:
                from quiz_url_tracker import module_from_filename, upsert_quiz_link

                upsert_quiz_link(
                    module=module_from_filename(src.name),
                    filename=src.name,
                    kahoot_title=title,
                    creator_url=creator_url,
                    share_url=share_url,
                    status="Uploaded",
                    path=TRACKER_PATH,
                )
                log(f"Updated share tracker: {TRACKER_PATH.name}")
            except Exception as tracker_exc:
                log(f"WARN: could not update share tracker: {tracker_exc}")
            successes += 1
            log(f"[{index}/{len(prepared)}] SUCCESS: {title}")
            if creator_url:
                log(f"Creator URL: {creator_url}")
            if share_url and share_url != creator_url:
                log(f"Share URL: {share_url}")
        except KeyboardInterrupt:
            append_result(src.name, title, "STOPPED", "Interrupted by user")
            log("Stopped by user.")
            break
        except Exception as exc:
            screenshot, html = capture_failure(page, src)
            details = f"{type(exc).__name__}: {exc}; screenshot={screenshot}; html={html}"
            append_result(src.name, title, "FAILED", details)
            try:
                from quiz_url_tracker import module_from_filename, upsert_quiz_link

                upsert_quiz_link(
                    module=module_from_filename(src.name),
                    filename=src.name,
                    kahoot_title=title,
                    status="Failed",
                    notes=str(exc)[:200],
                    path=TRACKER_PATH,
                )
            except Exception:
                pass
            log(f"[{index}/{len(prepared)}] FAILED: {exc}")
            log("Continuing with the next spreadsheet.")
        page.wait_for_timeout(1_500)

    log(f"Finished. Successful uploads: {successes}/{len(prepared)}")
    log(f"Results: {RESULTS_FILE}")
    log(f"Share links workbook: {TRACKER_PATH}")
    context.close()
    return 0 if successes == len(prepared) else 2


def main() -> int:
    args = parse_args()
    if args.prepare_only:
        return run(None, args)
    try:
        with sync_playwright() as playwright:
            return run(playwright, args)
    except PlaywrightTimeoutError as exc:
        print(f"Playwright timed out: {exc}", file=sys.stderr)
        return 3
    except TimeoutError as exc:
        print(f"{exc}", file=sys.stderr)
        return 4


if __name__ == "__main__":
    raise SystemExit(main())
