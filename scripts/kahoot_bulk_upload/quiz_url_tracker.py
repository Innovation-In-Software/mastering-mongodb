#!/usr/bin/env python3
"""Maintain kahoot/Kahoot_Quiz_Share_Links.xlsx for MD287 quiz URLs."""

from __future__ import annotations

import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[2]
SCRIPTS = REPO / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))
from mongodb_kahoot_questions import MODULE_TITLES, day_for_module  # noqa: E402

TRACKER_PATH = REPO / "kahoot" / "Kahoot_Quiz_Share_Links.xlsx"
SHEET_NAME = "Quiz Share Links"

HEADERS = [
    "Module",
    "Day",
    "Title",
    "Kahoot Title",
    "Source File",
    "Creator URL",
    "Share / Play URL",
    "Status",
    "Last Updated",
    "Notes",
]

MAX_MODULE = 8
TITLE_PREFIX = "Mastering MongoDB — "


def source_file_for_module(module: int) -> str:
    return f"kahoot/Kahoot_Module_{module}.xlsx"


def _style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for col, _ in enumerate(HEADERS, start=1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _autosize(ws) -> None:
    widths = {1: 10, 2: 8, 3: 52, 4: 64, 5: 32, 6: 55, 7: 55, 8: 14, 9: 20, 10: 36}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _read_existing(path: Path) -> dict[int, dict[str, str]]:
    existing: dict[int, dict[str, str]] = {}
    if not path.exists():
        return existing
    wb_old = load_workbook(path)
    ws_old = (
        wb_old["Quiz Share Links"]
        if "Quiz Share Links" in wb_old.sheetnames
        else wb_old.active
    )
    headers = [c.value for c in ws_old[1]]
    if not headers or headers[0] != "Module":
        return existing
    for row in ws_old.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        try:
            module = int(row[0])
        except (TypeError, ValueError):
            continue
        if module < 1 or module > MAX_MODULE:
            continue
        existing[module] = {
            headers[i]: ("" if row[i] is None else str(row[i]))
            for i in range(min(len(headers), len(row)))
            if headers[i]
        }
    return existing


def ensure_tracker(path: Path = TRACKER_PATH) -> Path:
    """Create the tracker if missing; refresh rows while preserving URLs/status."""
    path.parent.mkdir(parents=True, exist_ok=True)
    existing = _read_existing(path)

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    _style_header(ws)

    for module in range(1, MAX_MODULE + 1):
        title = MODULE_TITLES[module]
        prev = existing.get(module, {})
        creator = prev.get("Creator URL") or ""
        share = prev.get("Share / Play URL") or ""
        status = prev.get("Status") or "Pending"
        if (creator or share) and status == "Pending":
            status = "Uploaded"
        kahoot_title = f"{TITLE_PREFIX}Module {module} — {title}"
        ws.append(
            [
                module,
                day_for_module(module),
                title,
                kahoot_title,
                source_file_for_module(module),
                creator or None,
                share or None,
                status,
                prev.get("Last Updated") or None,
                prev.get("Notes") or None,
            ]
        )

    info = wb.create_sheet("How to use")
    info["A1"] = "Kahoot quiz share links — Mastering MongoDB"
    info["A1"].font = Font(bold=True, size=14)
    lines = [
        "",
        f"Purpose: one place to copy Creator / Share URLs for Modules 1–{MAX_MODULE}.",
        "",
        "Columns:",
        "- Creator URL: open/edit link from create.kahoot.it (auto-filled by bulk uploader when possible).",
        "- Share / Play URL: link you give facilitators or post in chat (paste from Kahoot Share if blank).",
        "- Status: Pending / Uploaded / Failed.",
        "",
        "How URLs get filled:",
        "1) Prefer: run scripts/kahoot_bulk_upload — it updates this file after each successful save.",
        "2) Or paste URLs manually from Kahoot (Share button → Copy link).",
        "",
        "Tip: keep Share / Play URL as the one students/facilitators should open.",
        "Players join at https://kahoot.it with the classroom PIN.",
        f"File location: {path.as_posix()}",
    ]
    for i, line in enumerate(lines, start=2):
        info[f"A{i}"] = line
    info.column_dimensions["A"].width = 110

    _autosize(ws)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def module_from_filename(filename: str) -> int | None:
    match = re.search(r"Module[_\s-]?(\d+)", filename, re.I)
    return int(match.group(1)) if match else None


def upsert_quiz_link(
    *,
    module: int | None,
    filename: str,
    kahoot_title: str,
    creator_url: str = "",
    share_url: str = "",
    status: str = "Uploaded",
    notes: str = "",
    path: Path = TRACKER_PATH,
) -> Path:
    """Update one module row. Does not recreate the whole workbook."""
    if not path.exists():
        ensure_tracker(path)

    try:
        wb = load_workbook(path)
    except PermissionError:
        fallback = path.with_name(path.stem + "_autosave.xlsx")
        if fallback.exists():
            wb = load_workbook(fallback)
            path = fallback
        else:
            ensure_tracker(path)
            wb = load_workbook(path)

    ws = wb["Quiz Share Links"] if "Quiz Share Links" in wb.sheetnames else wb.active
    if module is None:
        module = module_from_filename(filename)
    if module is None:
        raise ValueError(f"Cannot determine module number from {filename}")

    target_row = None
    for row in range(2, ws.max_row + 1):
        val = ws.cell(row, 1).value
        try:
            if int(val) == module:
                target_row = row
                break
        except (TypeError, ValueError):
            continue
    if target_row is None:
        raise ValueError(f"Module {module} not found in tracker")

    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws.cell(target_row, 4).value = kahoot_title
    ws.cell(target_row, 5).value = source_file_for_module(module)
    if creator_url:
        cell = ws.cell(target_row, 6)
        cell.value = creator_url
        cell.hyperlink = creator_url
        cell.font = Font(color="0563C1", underline="single")
    if share_url:
        cell = ws.cell(target_row, 7)
        cell.value = share_url
        cell.hyperlink = share_url
        cell.font = Font(color="0563C1", underline="single")
    ws.cell(target_row, 8).value = status
    ws.cell(target_row, 9).value = now
    if notes:
        ws.cell(target_row, 10).value = notes

    try:
        wb.save(path)
        return path
    except PermissionError:
        fallback = path.with_name(path.stem + "_autosave.xlsx")
        wb.save(fallback)
        return fallback
    except OSError:
        fallback = path.with_name(path.stem + "_autosave.xlsx")
        wb.save(fallback)
        try:
            fallback.replace(path)
            return path
        except Exception:
            return fallback


def main() -> None:
    path = ensure_tracker()
    print(f"Wrote {path}")


if __name__ == "__main__":
    main()
