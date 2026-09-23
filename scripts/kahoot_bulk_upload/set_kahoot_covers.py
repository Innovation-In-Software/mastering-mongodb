#!/usr/bin/env python3
"""Upload a cover image onto each existing Mastering MongoDB Kahoot."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright

TOOL = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOL))
from kahoot_bulk_upload import PROFILE_DIR, log  # noqa: E402

sys.path.insert(0, str(TOOL.parents[1] / "scripts"))
from mongodb_kahoot_questions import MODULE_TITLES  # noqa: E402

REPO = TOOL.parents[1]
TRACKER = REPO / "kahoot" / "Kahoot_Quiz_Share_Links.xlsx"
COVERS = TOOL / "covers"


def creator_urls() -> list[tuple[int, str]]:
    wb = load_workbook(TRACKER, read_only=True)
    ws = wb["Quiz Share Links"]
    found: list[tuple[int, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        module = int(row[0])
        url = str(row[5] or "").strip()
        if url.startswith("http"):
            found.append((module, url))
    return found


def dismiss_cookies(page) -> None:
    for pattern in [r"allow all", r"accept all", r"accept cookies", r"^agree$"]:
        try:
            btn = page.get_by_role("button", name=re.compile(pattern, re.I)).first
            if btn.is_visible(timeout=1200):
                btn.click()
                page.wait_for_timeout(800)
                log(f"Dismissed cookies via: {pattern}")
                return
        except Exception:
            continue


def open_settings(page) -> None:
    page.get_by_role("button", name=re.compile(r"^settings$", re.I)).first.click(timeout=15_000)
    page.wait_for_timeout(800)
    page.get_by_text(re.compile(r"^cover image$", re.I)).first.wait_for(state="visible", timeout=10_000)
    log("Opened Kahoot settings")


def _change_image_visible(page) -> bool:
    try:
        loc = page.get_by_text(re.compile(r"^change image$", re.I)).first
        loc.wait_for(state="visible", timeout=5_000)
        return True
    except Exception:
        return False


def finish_settings(page, module: int, title: str) -> None:
    try:
        portuguese = page.get_by_text(re.compile(r"^portuguese$", re.I)).first
        if portuguese.is_visible(timeout=1000):
            portuguese.click()
            page.get_by_role("option", name=re.compile(r"^english$", re.I)).first.click(timeout=5_000)
            page.wait_for_timeout(400)
            log("Set language to English")
    except Exception:
        pass
    blurb = f"Mastering MongoDB — Module {module}: {title}"
    try:
        box = page.locator("textarea").first
        if box.is_visible(timeout=1500):
            box.fill(blurb)
            log("Set description")
    except Exception:
        pass
    done = page.get_by_role("button", name=re.compile(r"^done$", re.I)).first
    if done.is_visible(timeout=2000):
        done.click(timeout=8_000)
        page.wait_for_timeout(1200)
        log("Closed settings")
    save = page.get_by_role("button", name=re.compile(r"^save$", re.I)).first
    try:
        if save.is_visible(timeout=2000) and save.is_enabled():
            save.click(timeout=8_000)
            page.wait_for_timeout(2000)
            log("Saved kahoot")
    except Exception:
        log("Save already completed")


def upload_cover(page, image: Path) -> None:
    if _change_image_visible(page):
        shot = TOOL / "logs" / "cover_already.png"
        page.screenshot(path=str(shot))
        log("Cover already on this kahoot")
        return
    add = page.locator(
        '[data-functional-selector="dialog-information-kahoot__image_library_btn"]'
    ).first
    add.wait_for(state="visible", timeout=8_000)
    chooser = None
    try:
        with page.expect_file_chooser(timeout=4_000) as info:
            add.click()
        chooser = info.value
        log("Cover button opened a file chooser")
    except PlaywrightTimeoutError:
        log("Cover button opened the image library")

    if chooser is not None:
        chooser.set_files(str(image))
    else:
        page.get_by_role("button", name=re.compile(r"^upload media$", re.I)).first.click(timeout=8_000)
        page.get_by_text(re.compile(r"^upload image$", re.I)).first.wait_for(state="visible", timeout=8_000)
        page.wait_for_timeout(600)
        log("Opened Upload image")
        inputs = page.locator('input[type="file"]')
        if inputs.count() > 0:
            inputs.last.set_input_files(str(image))
            log("Set cover through file input")
        else:
            with page.expect_file_chooser(timeout=8_000) as info:
                page.get_by_role("button", name=re.compile(r"^upload media$", re.I)).last.click()
            info.value.set_files(str(image))
            log("Set cover through the upload dialog")
    # Kahoot applies the file when settings close. Give the upload a moment.
    page.wait_for_timeout(2500)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--module", type=int, action="append", help="Only these module numbers")
    args = parser.parse_args()
    wanted = set(args.module or [])
    rows = [(m, u) for m, u in creator_urls() if not wanted or m in wanted]
    if not rows:
        print("No creator URLs in the share-links workbook.")
        return 1

    failures = 0
    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(PROFILE_DIR),
            channel="chrome",
            headless=False,
            viewport={"width": 1400, "height": 900},
            args=["--disable-blink-features=AutomationControlled"],
            ignore_default_args=["--enable-automation"],
        )
        page = context.pages[0] if context.pages else context.new_page()
        for index, (module, url) in enumerate(rows, start=1):
            image = COVERS / f"module_{module:02d}.png"
            log(f"[{index}/{len(rows)}] Module {module} <- {image.name}")
            try:
                page.goto(url, wait_until="domcontentloaded", timeout=90_000)
                page.wait_for_timeout(2000)
                dismiss_cookies(page)
                page.get_by_role("button", name=re.compile(r"^save$", re.I)).first.wait_for(
                    state="visible", timeout=25_000
                )
                open_settings(page)
                upload_cover(page, image)
                finish_settings(page, module, MODULE_TITLES[module])
                log(f"[{index}/{len(rows)}] Cover set for module {module}")
            except Exception as exc:
                failures += 1
                shot = TOOL / "logs" / f"cover_fail_module_{module}.png"
                shot.parent.mkdir(exist_ok=True)
                try:
                    page.screenshot(path=str(shot), full_page=True)
                except Exception:
                    pass
                log(f"[{index}/{len(rows)}] FAILED module {module}: {exc}")
            page.wait_for_timeout(800)
        context.close()
    return 0 if failures == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
