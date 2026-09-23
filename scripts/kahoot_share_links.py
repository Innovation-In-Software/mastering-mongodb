#!/usr/bin/env python3
"""Maintain kahoot/Kahoot_Quiz_Share_Links.xlsx. Existing URLs are kept."""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]
SCRIPTS = REPO / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from mongodb_kahoot_questions import MODULE_TITLES, day_for_module  # noqa: E402

TRACKER_PATH = REPO / "kahoot" / "Kahoot_Quiz_Share_Links.xlsx"
SHEET_NAME = "Quiz Share Links"
TITLE_PREFIX = "Mastering MongoDB — "
MAX_MODULE = 8

HEADERS = [
    "Module",
    "Day",
    "Title",
    "Kahoot Title",
    "Source File",
    "Creator URL",
    "Share / Play URL",
    "Status",
    "Last Updated",
    "Notes",
]


def _style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _autosize(ws) -> None:
    widths = {1: 10, 2: 8, 3: 52, 4: 64, 5: 32, 6: 55, 7: 55, 8: 14, 9: 20, 10: 36}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 22
    ws.auto_filter.ref = f"A1:J{MAX_MODULE + 1}"
    ws.freeze_panes = "A2"


def _read_existing(path: Path) -> dict[int, dict[str, str]]:
    existing: dict[int, dict[str, str]] = {}
    if not path.exists():
        return existing
    wb_old = load_workbook(path)
    ws_old = wb_old[SHEET_NAME] if SHEET_NAME in wb_old.sheetnames else wb_old.active
    headers = [c.value for c in ws_old[1]]
    if not headers or headers[0] != "Module":
        return existing
    index = {name: i for i, name in enumerate(headers) if name}
    for row in ws_old.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        try:
            module = int(row[0])
        except (TypeError, ValueError):
            continue
        def cell(name: str) -> str:
            pos = index.get(name)
            if pos is None or pos >= len(row) or row[pos] is None:
                return ""
            return str(row[pos])

        existing[module] = {
            "Creator URL": cell("Creator URL"),
            "Share / Play URL": cell("Share / Play URL"),
            "Status": cell("Status"),
            "Last Updated": cell("Last Updated"),
            "Notes": cell("Notes"),
        }
    return existing


def write_tracker(path: Path = TRACKER_PATH) -> Path:
    existing = _read_existing(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    for module in range(1, MAX_MODULE + 1):
        title = MODULE_TITLES[module]
        prior = existing.get(module, {})
        creator = prior.get("Creator URL", "")
        share = prior.get("Share / Play URL", "")
        status = prior.get("Status") or ("Linked" if share else "Not uploaded")
        updated = prior.get("Last Updated") or ""
        notes = prior.get("Notes") or "Import kahoot/Kahoot_Module_%d.xlsx" % module
        ws.append(
            [
                module,
                day_for_module(module),
                title,
                f"{TITLE_PREFIX}Module {module} — {title}",
                f"kahoot/Kahoot_Module_{module}.xlsx",
                creator,
                share,
                status,
                updated,
                notes,
            ]
        )
    _style_header(ws)
    _autosize(ws)
    wb.save(path)
    return path


def main() -> int:
    path = write_tracker()
    print(f"wrote {path.relative_to(REPO)} at {datetime.now():%Y-%m-%d %H:%M}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
