#!/usr/bin/env python3
"""Shared rendering helpers for MD287 content + diagram slides.

Chrome, type, and color match the Innovation In Software house system
extracted from the Java Software Engineer Bootcamp gold decks
(Modules 1, 2, 8, 9, 12): Leelawadee titles, Nirmala UI body, crimson/navy
palette, right-edge red/black rail, cream KEY TAKEAWAY bar, 14pt footer,
rotated 23pt page numbers.
"""
from __future__ import annotations

import difflib
import os
import re
import time
from pathlib import Path

from mdb_bullet_expansions import expand as expand_collapsed_bullet
import mdb_speaker_notes as SN

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------------------
# Brand constants (gold house system)
# ---------------------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

FOOTER_TEXT = "© 2026 by Innovation In Software Corporation"
COURSE_NAME = "MASTERING MONGODB"
TITLE_FONT = "Leelawadee"
BODY_FONT = "Nirmala UI"
CODE_FONT = "Consolas"

RED = RGBColor(0xD1, 0x28, 0x2E)
NAVY = RGBColor(0x00, 0x20, 0x60)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x80, 0x80, 0x80)
INK = RGBColor(0x00, 0x00, 0x00)
MUTED = RGBColor(0x59, 0x59, 0x59)
CARD_BG = RGBColor(0xF5, 0xF5, 0xF5)
CARD_LINE = RGBColor(0xE0, 0xE0, 0xE0)
TAKEAWAY_BG = RGBColor(0xFF, 0xF3, 0xE0)
TABLE_HEADER_BG = NAVY
GREEN = RGBColor(0x2E, 0x7D, 0x32)
PURPLE = RGBColor(0x5B, 0x3A, 0x9E)
ORANGE = RGBColor(0xC9, 0x69, 0x0C)
TEAL = RGBColor(0x0E, 0x7C, 0x7B)
BADGE_COLORS = (NAVY, GREEN, ORANGE, PURPLE, TEAL, RED)
TITLE_BADGE_COLORS = (PURPLE, NAVY, TEAL, GREEN)

RAIL_X = Inches(13.125)
RAIL_W = Inches(0.208)
RAIL_SPLIT_Y = Inches(1.5)

CONTENT_X = Inches(0.50)
CONTENT_W = Inches(12.30)
TITLE_X = Inches(0.60)
TITLE_W = Inches(12.00)

FOOTER_FONT_SIZE = 14
FOOTER_BOX = (Inches(0.667), Inches(7.098), Inches(5.00), Inches(0.31))
NUMBER_BOX = (Inches(12.492), Inches(6.807), Inches(0.72), Inches(0.50))
TAKEAWAY_Y = Inches(6.50)
TAKEAWAY_H = Inches(0.62)

MARGIN = Inches(0.30)
PICTURE_BOX = (MARGIN, MARGIN, Inches(12.733), Inches(6.40))

MARKER_PREFIX = "module-diagram:"
NUMBER_SHAPE_NAME = "Diagram Page Number"
BRANDING_SLIDE_NAME = "teksystems-branding"
BRANDING_NOTES = (
    "TEKsystems Global Services — Workforce Development. "
    "This is the client branding slide. Pause for introductions, then continue."
)

TAKEAWAY_TEXT = INK


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank_layout(prs: Presentation):
    return prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]


def _strip_placeholders(slide) -> None:
    for shape in list(slide.placeholders):
        shape._element.getparent().remove(shape._element)


def add_textbox(slide, left, top, width, height, *, anchor=MSO_ANCHOR.TOP, wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for attr in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, attr, 0)
    return box


def add_run(paragraph, text, *, size, bold=False, italic=False, color=INK, font=BODY_FONT):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    if font:
        run.font.name = font
    return run


def add_text(slide, left, top, width, height, text, *, size, bold=False, italic=False,
             color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=BODY_FONT):
    box = add_textbox(slide, left, top, width, height, anchor=anchor)
    lines = (text or "").split("\n") or [""]
    for i, line in enumerate(lines):
        p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
        p.alignment = align
        add_run(p, line, size=size, bold=bold, italic=italic, color=color, font=font)
    return box


def set_fill(shape, color, *, line_color=None, line_w=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_w or 0.75)


def add_rail_and_footer(slide, page_num: int | None = None) -> None:
    top_rail = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, RAIL_X, Emu(0), RAIL_W, RAIL_SPLIT_Y)
    set_fill(top_rail, RED)

    bottom_rail = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, RAIL_X, RAIL_SPLIT_Y, RAIL_W, Inches(6.0)
    )
    set_fill(bottom_rail, BLACK)

    if page_num is not None:
        footer = add_text(slide, *FOOTER_BOX, FOOTER_TEXT, size=FOOTER_FONT_SIZE, color=GRAY)
        footer.text_frame.word_wrap = False

        number_text = str(page_num)
        number_size = 23 if len(number_text) <= 2 else 17
        number = add_text(
            slide, *NUMBER_BOX, number_text, size=number_size, color=RED, bold=True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        number.text_frame.word_wrap = False
        number.rotation = 270
        number.name = NUMBER_SHAPE_NAME


SUBTITLE_Y = Inches(0.82)
SUBTITLE_W = Inches(12.30)
SUBTITLE_LINE_H = Inches(0.26)
CONTENT_TOP_MIN = Inches(1.28)
# A subtitle wrapping to two lines pushes content this far down.
CONTENT_TOP_MAX = Inches(1.46)


def subtitle_line_count(subtitle: str) -> int:
    """Lines a 17pt italic subtitle wraps to across SUBTITLE_W."""
    if not subtitle:
        return 0
    chars_per_line = max(20, int(int(SUBTITLE_W) / 914400 / 0.124))
    return max(1, -(-len(subtitle) // chars_per_line))


def title_line_count(header: str, size: float) -> int:
    """Lines a title wraps to at `size` across TITLE_W."""
    return max(1, -(-len(header or "") // title_chars_per_line(size)))


def title_bottom(title: str | None) -> int:
    """Bottom edge of the title block -- merged titles can run to three lines."""
    header = (title or "").upper()
    if not header:
        return int(Inches(0.84))
    size = title_font_size(header)
    line_h = int(Inches(size / 72 * 1.25))
    return int(Inches(0.14)) + title_line_count(header, size) * line_h


def content_top_for(subtitle: str | None, title: str | None = None) -> int:
    """Y where body content may start without colliding with title or subtitle."""
    top = int(CONTENT_TOP_MIN)
    lines = subtitle_line_count(subtitle)
    if lines:
        top = max(top, int(SUBTITLE_Y) + lines * int(SUBTITLE_LINE_H) + int(Inches(0.08)))
    else:
        top = max(top, title_bottom(title) + int(Inches(0.10)))
    return top


# Measured from rendered decks: uppercase Leelawadee runs ~0.0089 in per point
# per character across TITLE_W. Rounded up so the model never under-wraps.
TITLE_CHAR_W = 0.0095
TITLE_SIZES = (28, 25, 22, 20, 18, 17)
MAX_TITLE_LINES = 2


def title_chars_per_line(size: float) -> int:
    return max(12, int(int(TITLE_W) / 914400 / (TITLE_CHAR_W * size)))


def title_font_size(header: str) -> float:
    """Largest size that keeps the title within MAX_TITLE_LINES; floor 17pt."""
    for size in TITLE_SIZES:
        if title_line_count(header, size) <= MAX_TITLE_LINES:
            return size
    return TITLE_SIZES[-1]


def add_title_block(slide, *, title, subtitle=None, part_label=None):
    header = (title or "").upper()
    size = title_font_size(header)
    add_text(
        slide, TITLE_X, Inches(0.14), TITLE_W, Inches(0.70), header,
        size=size, color=RED, font=TITLE_FONT, anchor=MSO_ANCHOR.TOP,
    )
    if subtitle:
        lines = subtitle_line_count(subtitle)
        add_text(
            slide, TITLE_X, SUBTITLE_Y, SUBTITLE_W,
            max(int(Inches(0.42)), lines * int(SUBTITLE_LINE_H)), subtitle,
            size=17, italic=True, color=MUTED,
        )
    if part_label:
        add_text(
            slide, Inches(10.4), Inches(0.14), Inches(2.4), Inches(0.30),
            part_label, size=11, bold=True, color=MUTED, align=PP_ALIGN.RIGHT,
        )
    return content_top_for(subtitle, title)


def takeaway_height(text: str) -> int:
    """Bar height for this takeaway -- grows so long text is never clipped."""
    chars_per_line = max(40, int(int(Inches(11.60)) / 914400 / 0.088))
    lines = max(1, -(-(len(text or "") + 16) // chars_per_line))
    return max(int(TAKEAWAY_H), int(Inches(0.24)) * lines + int(Inches(0.22)))


def add_key_takeaway(slide, text: str, top=None) -> None:
    # Pinned to the slide bottom (above the copyright footer), matching gold.
    h = takeaway_height(text)
    y = int(TAKEAWAY_Y if top is None else top) - (h - int(TAKEAWAY_H))
    bar = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, CONTENT_X, y, Inches(12.00), h
    )
    set_fill(bar, TAKEAWAY_BG)
    try:
        bar.adjustments[0] = 0.08
    except Exception:
        pass
    box = add_textbox(
        slide, Inches(0.72), y, Inches(11.60), h, anchor=MSO_ANCHOR.MIDDLE,
    )
    p = box.text_frame.paragraphs[0]
    add_run(p, "KEY TAKEAWAY   ", size=13, bold=True, color=RED)
    add_run(p, text, size=13, color=INK)


def _icon_badge(slide, left, top, d, icon, bg, scale: float = 1.0):
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, d, d)
    set_fill(oval, bg)
    if icon:
        add_text(
            slide, left, top, d, d, icon, size=12 * scale, bold=True, color=WHITE,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=None,
        )


BODY_PT = 13
TABLE_PT = 12
MAX_BODY_PT = 16          # hard ceiling on body text
MAX_TEXT_SCALE = MAX_BODY_PT / BODY_PT


def _chars_per_inch(width, scale: float = 1.0):
    return max(12, int(int(width) / 914400 / (0.095 * scale)))


HEADING_GAP_IN = 0.11   # breathing room under a panel heading


def heading_gap(scale: float = 1.0) -> int:
    return int(Inches(HEADING_GAP_IN) * scale)


def heading_height(heading, width, scale: float = 1.0, badge: bool = True) -> int:
    """Height of a panel heading, allowing for the lines it wraps to."""
    one_line = int(Inches(0.34) * scale)
    if not heading:
        return one_line
    indent = (int(Inches(0.34) * scale) if badge else 0) + int(Inches(0.08))
    usable = max(int(Inches(1.00)), int(width) - indent)
    cpi = max(8, int(usable / 914400 / (0.088 * scale)))
    lines = max(1, -(-len(heading) // cpi))
    return max(one_line, lines * int(Inches(0.26) * scale) + int(Inches(0.08) * scale))


def estimate_panel_height(panel, width, scale: float = 1.0) -> int:
    """Content-sized height in EMU — never stretched to the takeaway bar."""
    heading = panel.get("heading", "")
    if "table" in panel:
        headers = panel["table"]["headers"]
        rows = panel["table"]["rows"]
        col_w = packed_table_widths(headers, rows, width, scale)
        return (heading_height(heading, width, scale, badge=False) + heading_gap(scale)
                + sum(_table_row_heights(headers, rows, col_w, scale)))
    heading_h = heading_height(heading, width, scale) + heading_gap(scale)
    bullets = panel.get("bullets") or []
    cpi = _chars_per_inch(width, scale)
    h = heading_h
    for _marker, text in bullets:
        lines = max(1, -(-(len(text) + 4) // cpi))
        h += int(Inches(0.24) * scale) * lines
    return h + int(Inches(0.06) * scale)


def packed_table_widths(headers, rows, max_width, scale: float = 1.0) -> list[int]:
    n = len(headers)
    max_lens = [len(str(headers[c])) for c in range(n)]
    for row in rows:
        for c, val in enumerate(row):
            if c < n:
                max_lens[c] = max(max_lens[c], len(str(val)))
    raw = [
        max(int(Inches(1.15 * scale)), min(int(Inches(4.4)), int(Inches((0.11 * L + 0.40) * scale))))
        for L in max_lens
    ]
    total = sum(raw)
    max_w = int(max_width)
    long_cells = any(L > 22 for L in max_lens)
    if long_cells:
        scale = max_w / max(total, 1)
        raw = [max(int(Inches(1.05)), int(w * scale)) for w in raw]
        raw[-1] += max_w - sum(raw)
    elif total > max_w:
        scale = max_w / total
        raw = [max(int(Inches(0.90)), int(w * scale)) for w in raw]
    return raw


def _cell_chars_per_line(col_width: int, scale: float = 1.0) -> int:
    """Characters that fit in a table cell, allowing for its 0.1in side margins."""
    usable = max(int(Inches(0.45)), int(col_width) - int(Inches(0.20)))
    return max(6, int(usable / 914400 / (0.078 * scale)))


def _table_row_heights(headers, rows, col_w, scale: float = 1.0) -> list[int]:
    unit = int(Inches(0.26) * scale)
    pad = int(Inches(0.10) * scale)  # cell top+bottom margins

    header_lines = 1
    for c, val in enumerate(headers):
        if c < len(col_w):
            cpi = _cell_chars_per_line(col_w[c], scale)
            header_lines = max(header_lines, max(1, -(-len(str(val)) // cpi)))
    heights = [max(int(Inches(0.30) * scale), unit * header_lines + pad)]

    # Every body row gets the same height -- ragged rows read as a broken table.
    widest = 1
    for row in rows:
        for c, val in enumerate(row):
            if c >= len(col_w):
                continue
            cpi = _cell_chars_per_line(col_w[c], scale)
            widest = max(widest, max(1, -(-len(str(val)) // cpi)))
    heights.extend([unit * widest + pad] * len(rows))
    return heights


def is_wide_table(panel) -> bool:
    """Panels the renderer hoists to full width instead of placing in a column."""
    if "table" not in panel:
        return False
    return any(len(str(c)) > 36 for row in panel["table"]["rows"] for c in row)


def partition_panels(panels: list) -> tuple[list, list]:
    """Split panels the way `content_slide` does: (hoisted full-width, grid)."""
    remaining = list(panels)
    if len(remaining) >= 2 and any(is_wide_table(p) for p in remaining):
        return (
            [p for p in remaining if is_wide_table(p)],
            [p for p in remaining if not is_wide_table(p)],
        )
    return [], remaining


def grid_columns(n: int) -> int:
    """Columns for n panels. Three reads well for 3, 5, 6; 4 stays a clean 2x2."""
    if n <= 1:
        return 1
    if n == 2 or n == 4:
        return 2
    return 3


def layout_height(panels: list, gap: int | None = None, scale: float = 1.0) -> int:
    """Height `content_slide` will actually use for these panels.

    This mirrors the renderer exactly -- full-width hoisted tables stacked
    first, then the column grid -- so the packer never underestimates.
    """
    gap = int(Inches(0.12)) if gap is None else gap
    hoisted, grid = partition_panels(panels)

    total = 0
    for panel in hoisted:
        total += estimate_panel_height(panel, int(CONTENT_W), scale) + gap

    n = len(grid)
    if n:
        cols = grid_columns(n)
        col_w = (int(CONTENT_W) - gap * (cols - 1)) // cols
        for i in range(0, n, cols):
            row = grid[i:i + cols]
            total += max(estimate_panel_height(p, col_w, scale) for p in row) + gap
    return total


ENUMERATOR = re.compile(r"^\s*(\d{1,2}[.)])\s+")


def split_enumerator(text: str) -> tuple[str, str]:
    """("1.", "rest") when the bullet text already starts with its own number."""
    match = ENUMERATOR.match(text or "")
    if match:
        return match.group(1), text[match.end():]
    return "", text


LEAD_IN = re.compile(r"^([^:]{2,60}):\s+(?=\S)")


def split_lead_in(text: str) -> tuple[str, str]:
    """("Build with Spring Boot", "rest") for bullets written as "Label: rest".

    Only a short label at the very start counts -- not a colon that turns up
    mid-sentence, and not a URL scheme (those have no space after the colon).
    """
    match = LEAD_IN.match(text or "")
    if not match:
        return "", text
    label = match.group(1).strip()
    if ". " in label or label.endswith("."):
        return "", text
    return label, text[match.end():]


def _add_bullet_text(paragraph, text: str, size) -> None:
    """Bullet body with its lead-in label set bold black."""
    label, rest = split_lead_in(text)
    if label:
        add_run(paragraph, f"{label}: ", size=size, bold=True, color=INK)
        add_run(paragraph, rest, size=size, color=INK)
    else:
        add_run(paragraph, text, size=size, color=INK)


def _is_standard_or_emoji_marker(marker: str) -> bool:
    """True for plain bullets and emoji-only markers (render as a standard bullet)."""
    m = (marker or "").strip()
    if m in ("▪", "-", "•", ""):
        return True
    cps = [ord(ch) for ch in m]
    joiners = {0x200D, 0xFE0E, 0xFE0F, 0x20E3}
    if not cps or cps[0] in joiners:
        return False

    def emoji_cp(cp: int) -> bool:
        return (
            cp in joiners
            or 0x1F3FB <= cp <= 0x1F3FF
            or 0x1F000 <= cp <= 0x1FAFF
            or 0x2600 <= cp <= 0x27BF
            or 0x2300 <= cp <= 0x23FF
            or 0x2B50 <= cp <= 0x2B55
            or 0x1F1E6 <= cp <= 0x1F1FF
        )

    return all(emoji_cp(cp) for cp in cps)


def add_panel_bullets(slide, left, top, width, *, icon, heading, bullets,
                      badge_color=RED, scale: float = 1.0) -> int:
    """Heading + bullets, no surrounding card. Returns used height in EMU."""
    body_pt = BODY_PT * scale
    d = int(Inches(0.34) * scale)
    head_h = heading_height(heading, width, scale)
    _icon_badge(slide, left, top, d, icon, badge_color, scale=scale)
    add_text(
        slide,
        int(left) + d + int(Inches(0.08)),
        top,
        int(width) - d - int(Inches(0.08)),
        head_h,
        heading,
        size=body_pt, bold=True, color=RED, anchor=MSO_ANCHOR.TOP,
    )
    body_top = int(top) + head_h + heading_gap(scale)
    est = estimate_panel_height({"bullets": bullets, "heading": heading}, width, scale)
    body_h = max(int(Inches(0.28) * scale), est - head_h - heading_gap(scale))
    body = add_textbox(slide, left, body_top, width, body_h)
    tf = body.text_frame
    tf.word_wrap = True
    first = True
    for marker, text in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(3 * scale)
        p.space_before = Pt(0)
        if _is_standard_or_emoji_marker(marker):
            number, rest = split_enumerator(text)
            if number:
                # Text carries its own "1." -- use it instead of a bullet glyph.
                add_run(p, f"{number} ", size=body_pt, bold=True, color=RED)
                _add_bullet_text(p, rest, body_pt)
            else:
                add_run(p, "•  ", size=body_pt, color=INK)
                _add_bullet_text(p, text, body_pt)
        else:
            add_run(p, f"{marker} ", size=body_pt, bold=True, color=RED)
            _add_bullet_text(p, text, body_pt)
    return est


def add_panel_table(slide, left, top, width, *, icon, heading, headers, rows,
                    scale: float = 1.0) -> int:
    """Compact table sized to text, not stretched. Returns used height in EMU."""
    y = int(top)
    if icon or heading:
        head_h = heading_height(heading, width, scale, badge=False)
        head = add_textbox(slide, left, y, width, head_h)
        p = head.text_frame.paragraphs[0]
        if icon:
            add_run(p, icon + "  ", size=BODY_PT * scale, bold=True, color=RED, font=None)
        add_run(p, heading, size=BODY_PT * scale, bold=True, color=RED)
        y += head_h + heading_gap(scale)

    col_w = packed_table_widths(headers, rows, width, scale)
    table_w = sum(col_w)
    n_rows = len(rows) + 1
    row_heights = _table_row_heights(headers, rows, col_w, scale)
    table_h = sum(row_heights)
    graphic_frame = slide.shapes.add_table(n_rows, len(headers), left, y, table_w, table_h)
    table = graphic_frame.table
    for c, w in enumerate(col_w):
        table.columns[c].width = w
    for r, h in enumerate(row_heights):
        table.rows[r].height = h
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_HEADER_BG
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for r in p.runs:
                r.font.size = Pt(TABLE_PT * scale)
                r.font.bold = True
                r.font.color.rgb = WHITE
                r.font.name = BODY_FONT
    for ridx, row in enumerate(rows, start=1):
        bg = WHITE if ridx % 2 else CARD_BG
        for c, val in enumerate(row):
            cell = table.cell(ridx, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(TABLE_PT * scale)
                    r.font.bold = c == 0
                    r.font.color.rgb = INK
                    r.font.name = BODY_FONT
    return (y - int(top)) + table_h


def set_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def _panel_weight(panel, col_w) -> int:
    return estimate_panel_height(panel, col_w)


def fit_text_scale(panels: list, available: int) -> float:
    """Largest body-type scale (1.0..MAX_TEXT_SCALE) whose layout still fits."""
    if layout_height(panels, scale=MAX_TEXT_SCALE) <= available:
        return MAX_TEXT_SCALE
    lo, hi = 1.0, MAX_TEXT_SCALE
    for _ in range(18):  # bisect; scaling is monotonic in height
        mid = (lo + hi) / 2
        if layout_height(panels, scale=mid) <= available:
            lo = mid
        else:
            hi = mid
    return lo


MAX_PANELS_PER_SLIDE = 6

# 0.0 pins content under the title, 0.5 centres it in the free space.
CONTENT_VERTICAL_EASE = 0.5


def _balanced_partition(panels: list, k: int, usable: int, max_panels: int):
    """Split `panels` into k contiguous groups minimising the tallest group.

    Returns the groups, or None if no such split keeps every group within
    `usable` and at most `max_panels` panels.
    """
    n = len(panels)
    if k > n:
        return None

    # cost[i][j] = rendered height of panels[i:j]
    cost = [[None] * (n + 1) for _ in range(n + 1)]
    for i in range(n):
        for j in range(i + 1, min(n, i + max_panels) + 1):
            cost[i][j] = layout_height(panels[i:j])

    INF = float("inf")
    best = [[INF] * (k + 1) for _ in range(n + 1)]
    cut = [[-1] * (k + 1) for _ in range(n + 1)]
    best[0][0] = 0.0
    for t in range(1, k + 1):
        for j in range(1, n + 1):
            for i in range(max(0, j - max_panels), j):
                if best[i][t - 1] == INF or cost[i][j] is None:
                    continue
                worst = max(best[i][t - 1], cost[i][j])
                if worst < best[j][t]:
                    best[j][t] = worst
                    cut[j][t] = i
    if best[n][k] > usable:
        return None

    groups, j = [], n
    for t in range(k, 0, -1):
        i = cut[j][t]
        groups.append(panels[i:j])
        j = i
    return list(reversed(groups))


def available_body_height(subtitle: str | None = None, takeaway: str | None = None,
                          title: str | None = None) -> int:
    """Vertical room for body content, after title, subtitle and takeaway bar."""
    bar_growth = takeaway_height(takeaway or "") - int(TAKEAWAY_H)
    return int(TAKEAWAY_Y) - bar_growth - content_top_for(subtitle, title) - int(Inches(0.08))


def expand_bullets(panel: dict) -> dict:
    """Replace collapsed bullet lists with one explained point each."""
    bullets = panel.get("bullets")
    if not bullets:
        return panel
    out = []
    changed = False
    for marker, text in bullets:
        replacement = expand_collapsed_bullet(text)
        if replacement:
            out.extend((marker, line) for line in replacement)
            changed = True
        else:
            out.append((marker, text))
    if not changed:
        return panel
    expanded = dict(panel)
    expanded["bullets"] = out
    return expanded


def expanded_parts(slide_parts: list) -> list:
    return [[expand_bullets(p) for p in part] for part in slide_parts]


def pack_topic_panels(slide_parts: list, subtitle: str | None = None,
                      takeaway: str | None = None) -> list[list]:
    """Merge same-topic PART slides, filling each slide as fully as it can.

    Uses the fewest slides the content fits on, then balances panels across
    them -- so a topic never ends on a near-empty slide just because an earlier
    one hit a panel cap.
    """
    panels = [expand_bullets(p) for part in slide_parts for p in part]
    if not panels:
        return []

    usable = available_body_height(subtitle, takeaway)

    for k in range(1, len(panels) + 1):
        groups = _balanced_partition(panels, k, usable, MAX_PANELS_PER_SLIDE)
        if groups is not None:
            return groups
    return [[panel] for panel in panels]


def content_slide(prs, layout, *, page_num, title, subtitle, part_label, panels,
                   key_takeaway=None, notes=None):
    """Packed layout: no empty cards, tables sized to text, takeaway pinned to slide bottom."""
    slide = prs.slides.add_slide(layout)
    _strip_placeholders(slide)
    origin_y = int(
        add_title_block(slide, title=title, subtitle=subtitle, part_label=part_label)
    )
    gap = int(Inches(0.12))
    # A merged slide joins several takeaways, so the bar grows upward and the
    # body must yield that space too.
    bar_growth = (
        takeaway_height(key_takeaway) - int(TAKEAWAY_H) if key_takeaway else 0
    )
    available = int(TAKEAWAY_Y) - bar_growth - origin_y

    # Grow the body type until the panels fill the slide. A topic with three
    # bullets should not sit in 13pt over four inches of empty paper.
    scale = fit_text_scale(panels, available)

    # Whatever slack is left after scaling gets split above and below.
    slack = max(0, available - layout_height(panels, scale=scale))
    origin_y += int(slack * CONTENT_VERTICAL_EASE)

    hoisted, remaining = partition_panels(panels)
    y = origin_y
    badge_i = 0
    for panel in hoisted:
        used = add_panel_table(
            slide, int(CONTENT_X), y, int(CONTENT_W),
            icon=panel.get("icon", ""), heading=panel.get("heading", ""),
            scale=scale, **panel["table"],
        )
        y += used + gap
        badge_i += 1

    n = len(remaining)
    cols = grid_columns(n)
    col_w = (int(CONTENT_W) - gap * (max(cols, 1) - 1)) // max(cols, 1)
    i = 0
    while i < n:
        row = remaining[i:i + cols]
        row_h = 0
        x = int(CONTENT_X)
        for j, panel in enumerate(row):
            badge = BADGE_COLORS[(badge_i + j) % len(BADGE_COLORS)]
            if "table" in panel:
                used = add_panel_table(
                    slide, x, y, col_w, icon=panel.get("icon", ""),
                    heading=panel.get("heading", ""), scale=scale, **panel["table"],
                )
            else:
                used = add_panel_bullets(
                    slide, x, y, col_w, icon=panel.get("icon", ""),
                    heading=panel.get("heading", ""), bullets=panel["bullets"],
                    badge_color=badge, scale=scale,
                )
            row_h = max(row_h, used)
            x += col_w + gap
        y += row_h + gap
        badge_i += len(row)
        i += cols

    if key_takeaway:
        add_key_takeaway(slide, key_takeaway)
    add_rail_and_footer(slide, page_num)
    if notes:
        set_notes(slide, notes)
    return slide


# Largest area a diagram may use without touching the house chrome: clear of
# the right-edge rail and the rotated page number, and above the footer line.
DIAGRAM_LEFT = Inches(0.10)
DIAGRAM_TOP = Inches(0.08)
DIAGRAM_RIGHT = Inches(12.45)      # page number starts at 12.49in
DIAGRAM_BOTTOM = Inches(7.06)      # footer text starts at 7.10in


def add_picture_fitted(slide, path: Path):
    """Diagram as large as the slide allows -- never stretched, never cropped.

    Scaled uniformly to the limiting dimension of the free area, then centred in
    the slide space left of the rail (nudged left only if it would reach the
    page number).
    """
    with Image.open(path) as image:
        image_width, image_height = image.size
    box_w = int(DIAGRAM_RIGHT) - int(DIAGRAM_LEFT)
    box_h = int(DIAGRAM_BOTTOM) - int(DIAGRAM_TOP)
    scale = min(box_w / image_width, box_h / image_height)
    width = max(1, int(image_width * scale))
    height = max(1, int(image_height * scale))
    x = min((int(RAIL_X) - width) // 2, int(DIAGRAM_RIGHT) - width)
    x = max(int(DIAGRAM_LEFT), x)
    y = int(DIAGRAM_TOP) + (box_h - height) // 2
    return slide.shapes.add_picture(str(path), x, y, width=width, height=height)


def diagram_slide(prs, layout, *, page_num, path: Path, topic: str, notes: str | None = None):
    slide = prs.slides.add_slide(layout)
    slide._element.cSld.set("name", f"{MARKER_PREFIX}{path.name}")
    _strip_placeholders(slide)
    add_picture_fitted(slide, path)
    add_rail_and_footer(slide, page_num)
    if notes:
        set_notes(slide, notes)
    return slide


def _chapter_slide(prs, layout, *, tag, title, subtitle, quote=None, icons=None, module_label=None):
    slide = prs.slides.add_slide(layout)
    _strip_placeholders(slide)
    add_rail_and_footer(slide, page_num=None)

    if quote:
        add_text(
            slide, Inches(8.14), Inches(0.12), Inches(3.90), Inches(0.40), quote,
            size=14, italic=True, color=MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
        )
    add_text(
        slide, CONTENT_X, Inches(0.42), Inches(11.30), Inches(0.60), tag,
        size=23, color=NAVY, font=TITLE_FONT, anchor=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide, CONTENT_X, Inches(2.15), Inches(12.30), Inches(1.90), title,
        size=40, bold=True, color=RED, font=TITLE_FONT,
    )
    if subtitle:
        add_text(
            slide, CONTENT_X, Inches(4.25), Inches(11.00), Inches(0.85), subtitle,
            size=18, color=NAVY,
        )
    if icons:
        for i, ic in enumerate(icons[:4]):
            if isinstance(ic, dict):
                glyph, bg = ic.get("icon", ""), ic.get("color", TITLE_BADGE_COLORS[i % len(TITLE_BADGE_COLORS)])
            else:
                glyph, bg = ic, TITLE_BADGE_COLORS[i % len(TITLE_BADGE_COLORS)]
            cx = int(CONTENT_X) + i * int(Inches(0.85))
            oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx, Inches(5.30), Inches(0.70), Inches(0.70))
            set_fill(oval, bg)
            if glyph:
                add_text(
                    slide, cx, Inches(5.30), Inches(0.70), Inches(0.70), glyph,
                    size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                    anchor=MSO_ANCHOR.MIDDLE, font=None,
                )
    add_text(
        slide, CONTENT_X, Inches(6.75), Inches(8.23), Inches(0.40),
        f"{COURSE_NAME}  •  {module_label or tag}",
        size=14, color=MUTED, anchor=MSO_ANCHOR.MIDDLE,
    )
    return slide


def cover_slide(prs, layout, *, title, subtitle, tag, page_num=None, quote=None, icons=None):
    # page_num kept for call-site compatibility; title slides are unnumbered.
    return _chapter_slide(
        prs, layout, tag=tag, title=title, subtitle=subtitle,
        quote=quote, icons=icons, module_label=tag,
    )


def branding_image_path(root: Path) -> Path:
    jpg = root / "assets" / "design" / "teksystems-branding.jpg"
    png = root / "assets" / "design" / "teksystems-branding.png"
    if jpg.exists():
        return jpg
    if png.exists():
        return png
    raise FileNotFoundError(
        "TEKsystems branding image not found. Run scripts/create_branding_deck.py first."
    )


def branding_slide(prs, layout, image_or_root: Path):
    """Full-bleed TEKsystems branding slide (unnumbered), shown as slide 2."""
    path = image_or_root
    if path.is_dir():
        path = branding_image_path(path)
    slide = prs.slides.add_slide(layout)
    slide._element.cSld.set("name", BRANDING_SLIDE_NAME)
    _strip_placeholders(slide)
    slide.shapes.add_picture(str(path), Emu(0), Emu(0), SLIDE_W, SLIDE_H)
    set_notes(slide, BRANDING_NOTES)
    return slide


def closing_slide(prs, layout, *, title, subtitle, tag, quote=None, icons=None):
    complete = tag if tag.endswith("COMPLETE") else f"{tag} COMPLETE"
    return _chapter_slide(
        prs, layout, tag=complete, title=title, subtitle=subtitle,
        quote=quote, icons=icons, module_label=tag.replace(" COMPLETE", ""),
    )


def save_presentation(prs, out_path: Path, attempts: int = 6) -> None:
    """Save, retrying briefly.

    On Windows a sync client or indexer can hold a freshly written .pptx for a
    moment, which surfaces as OSError(EINVAL) from the zip writer.
    """
    out_path.parent.mkdir(parents=True, exist_ok=True)
    for attempt in range(1, attempts + 1):
        try:
            prs.save(str(out_path))
            return
        except OSError:
            if attempt == attempts:
                raise
            time.sleep(0.8 * attempt)


def build_standard_module(
    *,
    out_path: Path,
    topics: list,
    module_tag: str,
    module_number: int,
    title: str,
    subtitle: str,
    root: Path,
    quote: str | None = None,
    icons=None,
    closing_subtitle: str | None = None,
):
    """Cover (unnumbered) → TEKsystems branding → content slides with day diagrams → closing.

    Diagrams come from `curriculum/day-NN/module-MM/slide_diagrams/` — the
    day-level teaching set for that module. Content headers use the short
    `diagram_title`, not the long thesis `deck_title`.
    """
    prs = new_presentation()
    layout = blank_layout(prs)

    cover_slide(
        prs, layout, title=title, subtitle=subtitle, tag=module_tag,
        quote=quote, icons=icons,
    )
    branding_slide(prs, layout, root)

    diagrams = day_diagrams_for_module(root, module_number)
    placement = plan_diagram_placement(topics, diagrams)

    page = 1
    for index, topic in enumerate(topics):
        packed = pack_topic_panels(topic["slides"], topic["subtitle"])
        n_slides = len(packed)
        for i, panels in enumerate(packed, start=1):
            part_label = f"PART {i} OF {n_slides}" if n_slides > 1 else None
            content_slide(
                prs, layout, page_num=page,
                title=topic["diagram_title"],
                subtitle=topic["subtitle"],
                part_label=part_label, panels=panels,
                key_takeaway=topic["key_takeaway"],
            )
            page += 1

        diagram = placement.get(index)
        if diagram is not None:
            diagram_slide(
                prs, layout, page_num=page, path=diagram,
                topic=diagram.stem.split(" - ", 1)[-1],
            )
            page += 1

    closing_slide(
        prs, layout, title=title,
        subtitle=closing_subtitle or subtitle,
        tag=module_tag, quote=quote, icons=icons,
    )

    save_presentation(prs, out_path)
    print(f"Wrote {len(prs.slides)} slides ({len(diagrams)} diagrams) -> {out_path.name}")
    return prs


# ---------------------------------------------------------------------------
# Day-scoped slide diagrams (curriculum/day-NN/module-MM/slide_diagrams)
# ---------------------------------------------------------------------------
# Teaching PNGs live beside each module. Each day outline
# (curriculum/day-NN/slides_outline.md) groups numbered slides under
# "### Module N — ..." headings; a trailing "### Lab" block belongs to the
# last module taught that day. Per-topic PNGs live in module-MM/diagrams/.

DAY_COUNT = 5


def _day_outline_sections(outline_path: Path) -> list[tuple[int | None, list[int]]]:
    """[(module_number|None, [slide numbers]), ...] in document order."""
    sections: list[tuple[int | None, list[int]]] = []
    current: list[int] | None = None
    for line in outline_path.read_text(encoding="utf-8").splitlines():
        heading = re.match(r"^###\s+(.*)$", line)
        if heading:
            module = re.match(r"^Module\s+(\d+)\b", heading.group(1).strip())
            current = []
            sections.append((int(module.group(1)) if module else None, current))
            continue
        item = re.match(r"^(\d+)\.\s+\S", line)
        if item and current is not None:
            current.append(int(item.group(1)))
    return sections


def day_diagrams_for_module(root: Path, module_number: int) -> list[Path]:
    """Ordered day PNGs owned by `module_number`, including its trailing Lab block."""
    for day in range(1, DAY_COUNT + 1):
        day_dir = root / "curriculum" / f"day-{day:02d}"
        outline = day_dir / "slides_outline.md"
        if not outline.exists():
            continue
        sections = _day_outline_sections(outline)
        if not any(mod == module_number for mod, _ in sections):
            continue

        owner = None
        numbers: list[int] = []
        for mod, slide_numbers in sections:
            if mod is not None:
                owner = mod
            if owner == module_number:
                numbers.extend(slide_numbers)

        diagrams_dir = day_dir / f"module-{module_number:02d}" / "slide_diagrams"
        paths = []
        for n in sorted(numbers):
            matches = sorted(diagrams_dir.glob(f"{n:02d} - *.png"))
            if len(matches) != 1:
                raise SystemExit(
                    f"Day {day} slide {n:02d}: expected 1 diagram, found {len(matches)}"
                )
            paths.append(matches[0])
        return paths
    raise SystemExit(f"No day outline lists Module {module_number}")


def _title_similarity(a: str, b: str) -> float:
    def norm(s: str) -> str:
        s = re.sub(r"\([^)]*\)", " ", s.lower())
        s = re.sub(r"[^a-z0-9 ]", " ", s)
        return " ".join(s.split())

    return difflib.SequenceMatcher(None, norm(a), norm(b)).ratio()


def plan_diagram_placement(topics: list, diagrams: list[Path]) -> dict[int, Path]:
    """Map diagram -> topic index, in order, one topic per diagram.

    Both sequences follow the same curriculum arc, so placement is a monotonic
    alignment: title similarity decides where a diagram clearly belongs, and an
    even-spread prior places the rest. Returns {topic_index: diagram_path}.
    """
    n_topics, n_diagrams = len(topics), len(diagrams)
    if not n_diagrams:
        return {}
    if n_diagrams > n_topics:
        raise SystemExit(
            f"{n_diagrams} diagrams cannot be placed across {n_topics} topics"
        )

    titles = [t["diagram_title"] for t in topics]

    def score(i: int, j: int) -> float:
        sim = _title_similarity(diagrams[i].stem.split(" - ", 1)[-1], titles[j])
        want = i / (n_diagrams - 1) if n_diagrams > 1 else 0.0
        got = j / (n_topics - 1) if n_topics > 1 else 0.0
        return 0.75 * sim + 0.25 * (1.0 - abs(want - got))

    NEG = float("-inf")
    # best[i][j] = best total score placing diagrams 0..i with diagram i at topic j
    best = [[NEG] * n_topics for _ in range(n_diagrams)]
    back = [[-1] * n_topics for _ in range(n_diagrams)]
    for j in range(n_topics):
        best[0][j] = score(0, j)
    for i in range(1, n_diagrams):
        running, arg = NEG, -1
        for j in range(i, n_topics):
            if best[i - 1][j - 1] > running:
                running, arg = best[i - 1][j - 1], j - 1
            if running > NEG:
                best[i][j] = running + score(i, j)
                back[i][j] = arg

    last = max(range(n_diagrams - 1, n_topics), key=lambda j: best[n_diagrams - 1][j])
    placement: dict[int, Path] = {}
    for i in range(n_diagrams - 1, -1, -1):
        placement[last] = diagrams[i]
        last = back[i][last]
    return placement


# ---------------------------------------------------------------------------
# Day-wise decks
# ---------------------------------------------------------------------------
# One deck per training day instead of one per module. A day deck runs:
# day cover -> [module divider -> content + diagrams] x N -> day closing, with
# page numbers continuous across the whole day.


def day_modules(root: Path, day: int) -> list[int]:
    """Module numbers taught on `day`, in teaching order."""
    outline = root / "curriculum" / f"day-{day:02d}" / "slides_outline.md"
    seen: list[int] = []
    for module, _ in _day_outline_sections(outline):
        if module is not None and module not in seen:
            seen.append(module)
    return seen


def module_curriculum_dir(root: Path, module_number: int) -> Path:
    """curriculum/day-NN/module-MM for a module number."""
    for day in range(1, DAY_COUNT + 1):
        if module_number in day_modules(root, day):
            return root / "curriculum" / f"day-{day:02d}" / f"module-{module_number:02d}"
    raise SystemExit(f"No day outline lists Module {module_number}")


def day_meta(root: Path, day: int) -> dict:
    """Theme, learning outcome, and lab title parsed from DAY-CURRICULUM.md."""
    text = (root / "curriculum" / f"day-{day:02d}" / "DAY-CURRICULUM.md").read_text(
        encoding="utf-8"
    )

    theme = re.search(r"^#\s*Day\s*\d+\s*Curriculum\s*[—-]\s*(.+?)\s*$", text, re.M)
    outcome = re.search(r"^##\s*Day learning outcome\s*\n+(.+?)\s*$", text, re.M)
    lab = re.search(r"^\|\s*(Lab\s*\d+)\s*\|\s*(.+?)\s*\|\s*$", text, re.M)
    if not (theme and outcome):
        raise SystemExit(f"Day {day}: could not parse theme/outcome from DAY-CURRICULUM.md")
    return {
        "theme": theme.group(1),
        "outcome": outcome.group(1),
        "lab": f"{lab.group(1)} — {lab.group(2)}" if lab else None,
    }


def _wrap_title(text: str, width: int = 34) -> str:
    """Break a day theme across two lines so it fits the 40pt cover type."""
    words, lines, line = text.split(), [], ""
    for word in words:
        trial = f"{line} {word}".strip()
        if len(trial) > width and line:
            lines.append(line)
            line = word
        else:
            line = trial
    if line:
        lines.append(line)
    return "\n".join(lines[:2]) if len(lines) <= 2 else "\n".join(
        [" ".join(lines[: len(lines) // 2]), " ".join(lines[len(lines) // 2:])]
    )


def _rel(root: Path, path: Path) -> str:
    try:
        return Path(path).resolve().relative_to(Path(root).resolve()).as_posix()
    except ValueError:
        return Path(path).as_posix()


def _panel_dump(panels: list) -> list[dict]:
    """Plain-data copy of panels for the slide inventory."""
    out = []
    for panel in panels:
        entry = {
            "heading": panel.get("heading", ""),
            "bullets": [text for _marker, text in panel.get("bullets") or []],
        }
        if panel.get("table"):
            entry["table"] = panel["table"]
        out.append(entry)
    return out


def _record(prs, slide, **info) -> None:
    SN.record(position=prs.slides.index(slide) + 1, **info)


def _chapter_notes(prs, slide, kind: str, key: str, title: str, subtitle: str = "") -> None:
    text = SN.notes_for(key)
    if text:
        set_notes(slide, text)
    _record(prs, slide, kind=kind, keys=[key], title=title, subtitle=subtitle)


def _render_topic_slides(prs, layout, page: int, topics: list) -> tuple[int, int]:
    """Render lab/exercise topics as content slides. Returns (next_page, slide_count)."""
    count = 0
    for topic in topics:
        number = str(topic["number"])
        if number.startswith("ex-"):
            kind, key = "exercise", SN.exercise_key(number)
        else:
            kind, key = "lab", SN.lab_key(number)
        packed = pack_topic_panels(
            topic["slides"], topic["subtitle"], topic["key_takeaway"]
        )
        n_slides = len(packed)
        for i, panels in enumerate(packed, start=1):
            slide = content_slide(
                prs, layout, page_num=page,
                title=topic["diagram_title"], subtitle=topic["subtitle"],
                part_label=f"PART {i} OF {n_slides}" if n_slides > 1 else None,
                panels=panels, key_takeaway=topic["key_takeaway"],
                notes=SN.notes_for(key, headings=[p.get("heading") for p in panels],
                                   part=i, parts=n_slides),
            )
            _record(prs, slide, kind=kind, keys=[key], title=topic["diagram_title"],
                    subtitle=topic["subtitle"], part=i, parts=n_slides,
                    panels=_panel_dump(panels), key_takeaway=topic["key_takeaway"])
            page += 1
            count += 1
    return page, count


def build_day_deck(
    *,
    out_path: Path,
    day: int,
    root: Path,
    module_meta: dict[int, dict],
    module_topics: dict[int, list],
    lab_topics: list | None = None,
    exercise_topics: dict[int, list] | None = None,
):
    """Build one deck for a whole training day: modules + checkpoints, then the day's lab."""
    meta = day_meta(root, day)
    modules = day_modules(root, day)
    day_tag = f"DAY {day}"
    day_title = _wrap_title(meta["theme"])
    day_icons = module_meta[modules[0]].get("icons")
    day_quote = f'"{meta["outcome"]}"'

    prs = new_presentation()
    layout = blank_layout(prs)

    subtitle = meta["outcome"]
    if meta["lab"]:
        subtitle = f"{subtitle}  ·  {meta['lab']}"
    slide = cover_slide(
        prs, layout, title=day_title, subtitle=subtitle, tag=day_tag,
        quote=day_quote, icons=day_icons,
    )
    _chapter_notes(prs, slide, "cover", SN.cover_key(day), meta["theme"], subtitle)
    slide = branding_slide(prs, layout, root)
    _chapter_notes(prs, slide, "branding", SN.BRANDING_KEY, "TEKsystems branding")

    # Decide up front which text slides are better shown as their diagram.
    day_plan = []
    for module in modules:
        topics = module_topics[module]
        placement = plan_diagram_placement(topics, day_diagrams_for_module(root, module))
        day_plan.append((module, topics, placement))
    swapped = plan_text_to_diagram_swaps(day_plan)
    plan = plan_day_slides(day_plan, swapped)

    placements = {module: placement for module, _, placement in day_plan}
    rendered_modules = set()
    page = 1
    day_diagram_total = swap_total = text_total = merged_total = 0
    exercise_slides = 0
    prev_module: int | None = None

    def flush_exercises(mod: int) -> None:
        nonlocal page, exercise_slides
        if not exercise_topics:
            return
        extra = exercise_topics.get(mod) or []
        if not extra:
            return
        page, n = _render_topic_slides(prs, layout, page, extra)
        exercise_slides += n

    for kind, module, payload in plan:
        if module not in rendered_modules:
            if prev_module is not None:
                flush_exercises(prev_module)
            info = module_meta[module]
            slide = _chapter_slide(
                prs, layout, tag=info["module_tag"], title=info["title"],
                subtitle=info["subtitle"], quote=info.get("quote"),
                icons=info.get("icons"),
                module_label=day_tag + "  " + chr(8226) + "  " + info["module_tag"],
            )
            _chapter_notes(prs, slide, "module", SN.module_key(module),
                           info["title"].replace(chr(10), " "), info["subtitle"])
            rendered_modules.add(module)
            prev_module = module
        topics = module_topics[module]

        if kind == "swap":
            topic = topics[payload]
            key = SN.topic_key(module, topic["number"])
            image = module_diagram_path(root, module, topic["number"])
            text = SN.notes_for(key, as_diagram=True)
            slide = diagram_slide(
                prs, layout, page_num=page, path=image,
                topic=topic["diagram_title"],
                notes=text,
            )
            _record(prs, slide, kind="swap", keys=[key], title=topic["diagram_title"],
                    subtitle=topic["subtitle"], image=_rel(root, image), module=module,
                    panels=_panel_dump(topic_panels(topic)),
                    key_takeaway=topic["key_takeaway"])
            page += 1
            swap_total += 1

        elif kind == "day":
            diagram = placements[module][payload]
            key = SN.day_diagram_key(root, diagram)
            slide = diagram_slide(
                prs, layout, page_num=page, path=diagram,
                topic=diagram.stem.split(" - ", 1)[-1],
                notes=SN.notes_for(key, as_diagram=True),
            )
            _record(prs, slide, kind="day-diagram", keys=[key],
                    title=diagram.stem.split(" - ", 1)[-1], image=_rel(root, diagram),
                    module=module)
            page += 1
            day_diagram_total += 1

        elif len(payload) == 1:
            topic = topics[payload[0]]
            packed = pack_topic_panels(
                topic["slides"], topic["subtitle"], topic["key_takeaway"]
            )
            n_slides = len(packed)
            key = SN.topic_key(module, topic["number"])
            for i, panels in enumerate(packed, start=1):
                slide = content_slide(
                    prs, layout, page_num=page,
                    title=topic["diagram_title"], subtitle=topic["subtitle"],
                    part_label=f"PART {i} OF {n_slides}" if n_slides > 1 else None,
                    panels=panels, key_takeaway=topic["key_takeaway"],
                    notes=SN.notes_for(key, headings=[p.get("heading") for p in panels],
                                       part=i, parts=n_slides),
                )
                _record(prs, slide, kind="topic", keys=[key], title=topic["diagram_title"],
                        subtitle=topic["subtitle"], part=i, parts=n_slides, module=module,
                        panels=_panel_dump(panels), key_takeaway=topic["key_takeaway"])
                page += 1
                text_total += 1

        else:
            # Several small topics share one slide.
            panels = [q for i in payload for q in topic_panels(topics[i])]
            keys = [SN.topic_key(module, topics[i]["number"]) for i in payload]
            legacy = [
                SN.legacy_topic_block(topics[i]["diagram_title"], topics[i]["subtitle"],
                                      topics[i]["key_takeaway"])
                for i in payload
            ]
            notes = SN.merged_notes(keys, [topics[i]["diagram_title"] for i in payload], legacy)
            slide = content_slide(
                prs, layout, page_num=page,
                title=merged_title(topics, payload), subtitle=None,
                part_label=None, panels=panels,
                key_takeaway=merged_takeaway(topics, payload), notes=notes,
            )
            _record(prs, slide, kind="merged", keys=keys, module=module,
                    title=merged_title(topics, payload),
                    panels=_panel_dump(panels),
                    key_takeaway=merged_takeaway(topics, payload),
                    topics=[
                        {"key": k, "title": topics[i]["diagram_title"],
                         "subtitle": topics[i]["subtitle"],
                         "key_takeaway": topics[i]["key_takeaway"],
                         "panels": _panel_dump(topic_panels(topics[i]))}
                        for k, i in zip(keys, payload)
                    ])
            page += 1
            text_total += 1
            merged_total += len(payload)

    if prev_module is not None:
        flush_exercises(prev_module)

    # Lab section: the day's hands-on work, from labs/day-NN/labN/LAB-N-GUIDE.md
    lab_slides = 0
    if lab_topics:
        lab_label = meta["lab"] or f"LAB {day}"
        slide = _chapter_slide(
            prs, layout, tag=f"LAB {day}", title=_wrap_title(lab_label),
            subtitle=meta["outcome"], quote=day_quote, icons=day_icons,
            module_label=day_tag + "  " + chr(8226) + "  LAB " + str(day),
        )
        _chapter_notes(prs, slide, "lab-divider", SN.lab_divider_key(day),
                       lab_label, meta["outcome"])
        page, lab_slides = _render_topic_slides(prs, layout, page, lab_topics)

    slide = closing_slide(
        prs, layout, title=day_title,
        subtitle="Day " + str(day) + " complete " + chr(8212) + " " + meta["outcome"],
        tag=day_tag, quote=day_quote, icons=day_icons,
    )
    _chapter_notes(prs, slide, "closing", SN.closing_key(day), meta["theme"], meta["outcome"])

    save_presentation(prs, out_path)
    missing = SN.finish_day(day)
    if missing:
        print(f"  speaker notes: {len(missing)} slide unit(s) without authored notes")
        if os.environ.get("MD287_REQUIRE_NOTES") == "1":
            raise SystemExit("missing speaker notes: " + ", ".join(missing[:25]))
    total_diagrams = day_diagram_total + swap_total
    pct = round(100 * total_diagrams / max(total_diagrams + text_total, 1))
    print(
        f"Wrote {len(prs.slides)} slides (modules {'+'.join(str(m) for m in modules)}) "
        f"-> {text_total} text / {total_diagrams} diagram "
        f"({day_diagram_total} day + {swap_total} swapped) = {pct}% visual; "
        f"{merged_total} topics merged, {exercise_slides} exercise slides, "
        f"{lab_slides} lab slides -> {out_path.name}"
    )
    return prs


# ---------------------------------------------------------------------------
# Text-slide -> module-diagram swaps (day decks only)
# ---------------------------------------------------------------------------
# The per-topic PNGs in curriculum/day-NN/module-MM/diagrams/ are what the
# deck text was written from, so each one is a superset of its text slide.
# Where the diagram explains the concept better, the day deck shows the
# diagram instead of the text -- bullets move into speaker notes.

VISUAL_TITLE_WORDS = (
    "architecture flow lifecycle pipeline topology layer stack sequence journey map anatomy "
    "structure overview pattern comparison stage phase process workflow hierarchy relationship "
    "communication interaction deployment rollout scaling model tree matrix strategy mesh "
    "gateway routing traffic boundary decomposition integration orchestration chain loop"
).split()


def module_diagram_path(root: Path, module: int, number) -> Path:
    """The per-topic PNG in curriculum/day-NN/module-MM/diagrams/."""
    folder = module_curriculum_dir(root, module) / "diagrams"
    matches = sorted(folder.glob(f"{number} - *.png")) or sorted(
        folder.glob(f"{int(number):02d} - *.png")
    )
    if len(matches) != 1:
        raise SystemExit(
            f"Module {module} topic {number}: expected 1 diagram, found {len(matches)}"
        )
    return matches[0]


def topic_text_metrics(topic) -> dict:
    panels = [p for part in topic["slides"] for p in part]
    bullets = [b for p in panels for _, b in p.get("bullets", [])]
    return {
        "parts": len(pack_topic_panels(topic["slides"], topic["subtitle"])),
        "panels": len(panels),
        "bullets": len(bullets),
        "chars": sum(len(b) for b in bullets),
        "tables": sum(1 for p in panels if "table" in p),
        "split": any(re.search(r"\(\d\s*/\s*\d\)", p.get("heading", "")) for p in panels),
    }


def score_topic_for_diagram(topic) -> float:
    """How much better the topic's diagram is than its text slide. Higher = swap."""
    m = topic_text_metrics(topic)
    title = topic["diagram_title"].lower()
    score = 0.0

    # A sparse text slide leaves the diagram carrying nearly all the content.
    chars = m["chars"]
    if chars < 120:
        score += 3.0
    elif chars < 200:
        score += 2.0
    elif chars < 300:
        score += 1.0
    elif chars > 600:
        score -= 1.5  # substantive talking points -- keep them on the slide

    if m["bullets"] == 0 and m["tables"]:
        score += 1.5           # bare table; the diagram gives it structure
    if m["parts"] > 1:
        score += 2.0           # text had to spill onto extra slides
    if m["split"]:
        score += 1.5           # table was chopped into (1/3), (2/3)...
    if m["panels"] >= 4:
        score += 1.0           # crowded slide

    if any(word in title for word in VISUAL_TITLE_WORDS):
        score += 2.0
    if re.search(r"\bvs\b|\bversus\b", title):
        score += 1.0
    return score


def swap_notes(topic) -> str:
    """Speaker notes for a swapped slide: the text that left the slide."""
    lines = [topic["subtitle"], ""]
    for part in expanded_parts(topic["slides"]):
        for panel in part:
            heading = panel.get("heading")
            if heading:
                lines.append(heading)
            for _marker, text in panel.get("bullets", []):
                lines.append(f"  - {text}")
            table = panel.get("table")
            if table:
                lines.append("  " + " | ".join(str(h) for h in table["headers"]))
                for row in table["rows"]:
                    lines.append("  " + " | ".join(str(c) for c in row))
            lines.append("")
    lines.append(f"KEY TAKEAWAY: {topic['key_takeaway']}")
    return "\n".join(lines).strip()


MAX_TOPICS_PER_SLIDE = 3
MAX_PANELS_PER_MERGED_TOPIC = 2
MAX_MERGED_TITLE_CHARS = 118


def merged_title(topics: list, idxs: list[int]) -> str:
    return "  ·  ".join(topics[i]["diagram_title"] for i in idxs)


def merged_takeaway(topics: list, idxs: list[int]) -> str:
    return "   •   ".join(topics[i]["key_takeaway"] for i in idxs)


def topic_panels(topic) -> list:
    return [expand_bullets(p) for part in topic["slides"] for p in part]


def group_fits(topics: list, idxs: list[int]) -> bool:
    """Can these adjacent topics share one slide?"""
    if len(idxs) <= 1:
        return True                      # a lone topic may still split into parts
    if len(idxs) > MAX_TOPICS_PER_SLIDE:
        return False
    if len(merged_title(topics, idxs)) > MAX_MERGED_TITLE_CHARS:
        return False
    panels = []
    for i in idxs:
        own = topic_panels(topics[i])
        if len(own) > MAX_PANELS_PER_MERGED_TOPIC:
            return False                 # keep busy topics on their own slide
        panels.extend(own)
    if len(panels) > MAX_PANELS_PER_SLIDE:
        return False
    return layout_height(panels) <= available_body_height(
        None, merged_takeaway(topics, idxs), merged_title(topics, idxs)
    )


def plan_day_slides(day_plan: list, swapped: set) -> list:
    """The day's slide sequence, in order.

    Items are ("text", module, [topic indices]), ("swap", module, index) for a
    topic shown as its own diagram, or ("day", module, index) for a curriculum
    diagram. Adjacent text topics merge onto one slide where they fit.
    """
    out = []
    for module, topics, placement in day_plan:
        group: list[int] = []

        def flush():
            if group:
                out.append(("text", module, list(group)))
                group.clear()

        for index in range(len(topics)):
            if (module, index) in swapped:
                flush()
                out.append(("swap", module, index))
            elif group_fits(topics, group + [index]):
                group.append(index)
            else:
                flush()
                group.append(index)
            if index in placement:
                flush()
                out.append(("day", module, index))
        flush()
    return out


def slide_kinds(plan: list, module_topics: dict) -> list[str]:
    """T/D sequence the plan renders to (a merged text group is one T)."""
    kinds = []
    for kind, module, payload in plan:
        if kind == "text":
            if len(payload) == 1:
                topic = module_topics[module][payload[0]]
                kinds.extend("T" * len(pack_topic_panels(
                    topic["slides"], topic["subtitle"], topic["key_takeaway"])))
            else:
                kinds.append("T")
        else:
            kinds.append("D")
    return kinds


def _max_diagram_run(kinds: list[str]) -> int:
    longest = run = 0
    for kind in kinds:
        run = run + 1 if kind == "D" else 0
        longest = max(longest, run)
    return longest


def plan_text_to_diagram_swaps(day_plan: list, *, max_run: int = 2) -> set[tuple[int, int]]:
    """Choose which topics show their diagram instead of their text.

    Counts against the *merged* slide plan, so balance reflects the slides that
    are actually produced. Greedy by score, stopping once diagram slides balance
    text slides, and never more than `max_run` diagrams back to back.
    """
    module_topics = {module: topics for module, topics, _ in day_plan}

    candidates = []
    for module, topics, placement in day_plan:
        for index, topic in enumerate(topics):
            if index in placement:
                continue  # its curriculum diagram already follows it
            candidates.append(((module, index), score_topic_for_diagram(topic)))
    candidates.sort(key=lambda c: -c[1])

    def counts(swapped):
        kinds = slide_kinds(plan_day_slides(day_plan, swapped), module_topics)
        return kinds.count("T"), kinds.count("D"), _max_diagram_run(kinds)

    swapped: set[tuple[int, int]] = set()
    text, diagrams, _ = counts(swapped)
    for key, _score in candidates:
        if diagrams >= text:
            break
        trial = swapped | {key}
        t, d, run = counts(trial)
        if run > max_run or d > t + 1:
            continue
        # A diagram slide holds one topic; a merged text slide holds two or
        # three. Never take a swap that costs the deck an extra slide.
        if t + d > text + diagrams:
            continue
        swapped, text, diagrams = trial, t, d
    return swapped
