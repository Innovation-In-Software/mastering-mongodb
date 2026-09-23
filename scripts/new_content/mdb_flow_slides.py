#!/usr/bin/env python3
"""Flow, code, exercise and diagram-with-panel slide layouts for the new decks.

Each public function returns a *step*: a callable ``step(prs, layout)`` that adds
one slide. ``build_sequence`` runs a list of steps (these steps and the original
``sNN(prs, layout)`` slide builders mix freely) and numbers every page by its real
position in the finished deck.

Layouts
- diagram   diagram image + side panel of key points (or points under a wide image)
- bridge    part opener: roadmap of parts, what we covered, the next question
- code      one to three code blocks with labels, plus short "why" points
- exercise  scenario, tasks and expected outcome (practice or official checkpoint)
- knowledge ten-question knowledge check (answers go in the speaker notes)
- compare   comparison table + side panel
- custom    free drawing inside the standard title/takeaway frame
"""
from __future__ import annotations

import math
from pathlib import Path

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

import mdb_deck_kit as K
import mdb_diagram_slides as DS
import mdb_speaker_notes as SN
import mdb_visuals as V
from mdb_visuals import (
    BLACK, CARD_BG, CARD_LINE, DARK_GRAY, GREEN, INK, LEFT, LIGHT_GRAY, LIGHT_GREEN,
    LIGHT_NAVY, LIGHT_PURPLE, LIGHT_RED, LIGHT_TEAL, MUTED, NAVY, ORANGE, PURPLE, RED, RIGHT,
    TEAL, WHITE, WIDTH,
)

LIGHT_ORANGE = V.RGBColor(0xFD, 0xF0, 0xE3)
LIGHT = {NAVY: LIGHT_NAVY, PURPLE: LIGHT_PURPLE, TEAL: LIGHT_TEAL, GREEN: LIGHT_GREEN,
         ORANGE: LIGHT_ORANGE, RED: LIGHT_RED, BLACK: LIGHT_GRAY, DARK_GRAY: LIGHT_GRAY}
BOTTOM = V.BODY_BOTTOM
NO_TAKEAWAY_BOTTOM = 6.92
GAP = 0.14
TEXT = 12


# ---------------------------------------------------------------------------
# Notes
# ---------------------------------------------------------------------------
def tp(point: str, explain: str) -> dict:
    return {"point": point, "explain": explain}


def merge_units(*units) -> dict:
    points, examples, scenario = [], [], None
    for unit in units:
        if not unit:
            continue
        if isinstance(unit, list):
            points += unit
            continue
        points += unit.get("talking_points", [])
        examples += unit.get("real_world_examples", []) or []
        scenario = scenario or unit.get("use_case_scenario")
    out: dict = {"talking_points": points}
    if examples:
        out["real_world_examples"] = examples
    if scenario:
        out["use_case_scenario"] = scenario
    return out


def notes_text(unit):
    if unit is None or isinstance(unit, str):
        return unit
    if isinstance(unit, list):
        unit = {"talking_points": unit}
    return SN.compose(unit)


# ---------------------------------------------------------------------------
# Sequence builder
# ---------------------------------------------------------------------------
def build_sequence(prs, layout, steps) -> None:
    base_footer = K.add_rail_and_footer

    def positional_footer(slide, page_num=None):
        base_footer(slide, len(prs.slides) if page_num is not None else None)

    K.add_rail_and_footer = positional_footer
    try:
        for step in steps:
            step(prs, layout)
    finally:
        K.add_rail_and_footer = base_footer


def _new(prs, layout, title, subtitle):
    n = len(prs.slides) + 1
    slide, top = V.content_slide(prs, layout, number=n, title=title, subtitle=subtitle)
    return slide, top, n


def _finish(slide, n, takeaway, notes):
    V.finish(slide, number=n, takeaway=takeaway, notes=notes_text(notes))


def _bottom(takeaway):
    return BOTTOM if takeaway else NO_TAKEAWAY_BOTTOM


# ---------------------------------------------------------------------------
# Panel items
#   ("Heading", "text")                         -> card
#   {"h": .., "t": .., "c": COLOR}              -> card
#   {"code": .., "label": .., "kind": good|bad|info, "size": 11}
#   {"callout": ("Label", "text")}
#   {"chips": [...], "h": .., "c": COLOR, "cols": 2}
#   {"checks": [...], "h": .., "c": COLOR, "mark": "✓"}
# ---------------------------------------------------------------------------
def _norm(item):
    if isinstance(item, tuple):
        return {"h": item[0], "t": item[1]}
    return item


def _kind(item):
    for key in ("code", "callout", "chips", "checks"):
        if key in item:
            return key
    return "card"


def _text_h(text, w, size, bold=False):
    return V._estimate_height(V._lines(text), w, size, bold)


def _check_rows_h(rows, w, size):
    return sum(max(0.34, _text_h(r, w - 0.42, size) + 0.08) for r in rows)


def item_height(item, w):
    item = _norm(item)
    kind = _kind(item)
    size = item.get("size", TEXT if kind != "code" else 11)
    if kind == "card":
        head = 0.40 if item.get("h") else 0.10
        return head + _text_h(item["t"], w - 0.30, size) + 0.16
    if kind == "code":
        lines = item["code"].split("\n")
        return (0.36 if item.get("label") else 0) + len(lines) * size * 1.30 / 72 + 0.34
    if kind == "callout":
        label, body = item["callout"]
        return _text_h(label.upper() + "   " + body, w - 0.40, size) + 0.32
    if kind == "chips":
        cols = item.get("cols", 2)
        rows = math.ceil(len(item["chips"]) / cols)
        return (0.38 if item.get("h") else 0) + rows * 0.40 + (rows - 1) * 0.08
    head = 0.38 if item.get("h") else 0
    return head + _check_rows_h(item["checks"], w, size)


def _heading(slide, x, y, w, text, color):
    V.text(slide, x, y, w, 0.34, text.upper(), size=13, bold=True, color=color)


def draw_item(slide, item, x, y, w, h):
    item = _norm(item)
    kind = _kind(item)
    color = item.get("c", NAVY)
    size = item.get("size", TEXT if kind != "code" else 11)
    if kind == "card":
        V.box(slide, x, y, w, h, "", fill=CARD_BG, line=CARD_LINE, radius=0.06)
        ty = y + 0.08
        if item.get("h"):
            V.text(slide, x + 0.15, y + 0.06, w - 0.30, 0.32, item["h"], size=13, bold=True,
                   color=color, anchor=MSO_ANCHOR.MIDDLE)
            ty = y + 0.40
        V.text(slide, x + 0.15, ty, w - 0.30, y + h - ty - 0.06, item["t"], size=size, color=INK)
    elif kind == "code":
        cy = y
        if item.get("label"):
            mark = {"good": ("✓", GREEN), "bad": ("✕", RED)}.get(item.get("kind"))
            lx = x
            if mark:
                V.badge(slide, x, y + 0.03, 0.28, mark[0], fill=mark[1], size=10)
                lx = x + 0.36
            label_color = mark[1] if mark else NAVY
            V.text(slide, lx, y, w - (lx - x), 0.34, item["label"], size=13, bold=True,
                   color=label_color, anchor=MSO_ANCHOR.MIDDLE)
            cy = y + 0.36
        fill = {"good": LIGHT_GREEN, "bad": LIGHT_RED}.get(item.get("kind"), CARD_BG)
        V.code(slide, x, cy, w, y + h - cy, item["code"], size=size, fill=fill)
    elif kind == "callout":
        label, body = item["callout"]
        V.callout(slide, x, y, w, h, label, body, size=size)
    elif kind == "chips":
        cy = y
        if item.get("h"):
            _heading(slide, x, y, w, item["h"], color)
            cy = y + 0.38
        V.chips(slide, x, cy, w, item["chips"], cols=item.get("cols", 2), h=0.40, gap=0.08,
                fill=LIGHT.get(color, LIGHT_NAVY), color=color, size=size)
    else:
        cy = y
        if item.get("h"):
            _heading(slide, x, y, w, item["h"], color)
            cy = y + 0.38
        mark = item.get("mark", "✓")
        for row in item["checks"]:
            rh = max(0.34, _text_h(row, w - 0.42, size) + 0.08)
            V.badge(slide, x, cy + 0.03, 0.28, mark, fill=color, size=10)
            V.text(slide, x + 0.40, cy, w - 0.40, rh, row, size=size, anchor=MSO_ANCHOR.MIDDLE)
            cy += rh


def stack_items(slide, x, y, w, h, items, gap=GAP):
    items = [_norm(i) for i in items]
    if not items:
        return
    heights = [item_height(i, w) for i in items]
    free = h - sum(heights) - gap * (len(items) - 1)
    grow = [k for k, i in enumerate(items) if _kind(i) in ("card", "callout")]
    if free > 0 and grow:
        extra = min(free / len(grow), 0.45)
        for k in grow:
            heights[k] += extra
    elif free < -0.02:
        V.WARNINGS.append(f"{V._current_slide_label}: panel needs {sum(heights):.2f}in, "
                          f"has {h:.2f}in")
    cy = y
    for item, ih in zip(items, heights):
        draw_item(slide, item, x, cy, w, ih)
        cy += ih + gap


def row_items(slide, x, y, w, h, items, gap=0.20):
    items = [_norm(i) for i in items]
    if not items:
        return
    iw = (w - gap * (len(items) - 1)) / len(items)
    need = max(item_height(i, iw) for i in items)
    ih = min(h, need + 0.30)
    if need > h + 0.02:
        V.WARNINGS.append(f"{V._current_slide_label}: row needs {need:.2f}in, has {h:.2f}in")
    for k, item in enumerate(items):
        draw_item(slide, item, x + k * (iw + gap), y, iw, ih)


PENDING_ASPECT = 16 / 9


def image_path(image_dir, n, name):
    """The diagram for slide ``n`` — or the file name it will have once it is generated."""
    matches = sorted(Path(image_dir).glob(f"{n:02d} - *.png"))
    return matches[0] if matches else Path(image_dir) / f"{n:02d} - {name}.png"


def _pending(slide, image, x, y, w, h):
    """A reserved frame for a diagram that is still being generated."""
    image = Path(image)
    box = V.box(slide, x, y, w, h, "", fill=CARD_BG, line=MUTED, line_w=1.5, dash=True,
                radius=0.03)
    box.name = f"Diagram placeholder {image.stem[:2]}"
    title = image.stem[5:] if image.stem[:2].isdigit() else image.stem
    cy = y + h / 2
    V.text(slide, x + 0.30, cy - 1.00, w - 0.60, 0.40, "DIAGRAM IN PROGRESS", size=13,
           bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    V.text(slide, x + 0.30, cy - 0.55, w - 0.60, 0.95, title, size=20, bold=True, color=NAVY,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    V.text(slide, x + 0.30, cy + 0.50, w - 0.60, 0.70,
           f"Space reserved — save “{image.name}” in {image.parent.name}/ and rebuild.",
           size=12, italic=True, color=MUTED, align=PP_ALIGN.CENTER)


def _picture(slide, image, x, y, w, h):
    pad = 0.08
    if not Path(image).exists():
        _pending(slide, image, x - pad, y - pad, w + 2 * pad, h + 2 * pad)
        return
    V.box(slide, x - pad, y - pad, w + 2 * pad, h + 2 * pad, "", fill=WHITE, line=CARD_LINE,
          radius=0.03)
    stream, _ = DS.trimmed_image(image)
    pic = slide.shapes.add_picture(stream, V.I(x), V.I(y), V.I(w), V.I(h))
    pic.name = f"Diagram {image.stem[:2]}"


# ---------------------------------------------------------------------------
# Slide steps
# ---------------------------------------------------------------------------
def diagram(image, title, subtitle, items=(), takeaway=None, notes=None, panel_w=3.85,
            mode=None):
    """Diagram image plus key points (side panel, or under a very wide image)."""
    def step(prs, layout):
        slide, top, n = _new(prs, layout, title, subtitle)
        aspect = (DS.trimmed_image(image)[1] if Path(image).exists() else PENDING_ASPECT)
        y0, y1 = top + 0.10, _bottom(takeaway)
        pad = 0.10
        use = mode or ("stack" if aspect >= 3.2 else "side")
        if not items:
            aw, ah = WIDTH - 2 * pad, y1 - y0 - 2 * pad
            iw = min(aw, ah * aspect)
            ih = iw / aspect
            _picture(slide, image, LEFT + pad + (aw - iw) / 2, y0 + pad + (ah - ih) / 2, iw, ih)
        elif use == "side":
            px = RIGHT - panel_w
            aw, ah = px - 0.30 - LEFT - 2 * pad, y1 - y0 - 2 * pad
            iw = min(aw, ah * aspect)
            ih = iw / aspect
            _picture(slide, image, LEFT + pad + (aw - iw) / 2, y0 + pad + (ah - ih) / 2, iw, ih)
            stack_items(slide, px, y0, panel_w, y1 - y0, items)
        else:
            # key points sit just above the takeaway; the diagram fills the space above them
            gap = 0.20
            cw = (WIDTH - gap * (len(items) - 1)) / len(items)
            rows_h = min(max(item_height(_norm(i), cw) for i in items) + 0.30,
                         (y1 - y0) * 0.45)
            room = y1 - y0 - rows_h - 0.30 - 2 * pad
            iw = WIDTH - 2 * pad
            ih = iw / aspect
            if ih > room:
                ih, iw = room, room * aspect
            ix, iy = LEFT + (WIDTH - iw) / 2, y0 + pad + (room - ih) / 2
            _picture(slide, image, ix, iy, iw, ih)
            row_items(slide, LEFT, y1 - rows_h, WIDTH, rows_h, items)
        _finish(slide, n, takeaway, notes)
    return step


def bridge(parts, current, title, subtitle, so_far, question, covers, notes=None,
           takeaway=None):
    """Part opener: roadmap of all parts, what we have covered, what comes next."""
    def step(prs, layout):
        slide, top, n = _new(prs, layout, title, subtitle)
        y0, y1 = top + 0.10, _bottom(takeaway)
        labels, fills = [], []
        for k, part in enumerate(parts):
            if k < current:
                labels.append(f"✓ {part}")
                fills.append(GREEN)
            elif k == current:
                labels.append(part)
                fills.append(RED)
            else:
                labels.append(part)
                fills.append(DARK_GRAY)
        V.chevrons(slide, LEFT, y0, WIDTH, 0.78, labels, fills, size=13)
        cy = y0 + 1.10
        lw = 5.40
        rows = so_far
        heights = [max(0.46, _text_h(row, lw - 0.80, 15) + 0.14) for row in rows]
        card_h = min(y1 - cy, max(2.30, 0.63 + sum(h + 0.06 for h in heights) + 0.25))
        body = V.card(slide, LEFT, cy, lw, card_h, "So far", head_fill=GREEN, head_h=0.46,
                      title_size=14)
        ry = body + 0.05
        for row in rows:
            rh = max(0.46, _text_h(row, lw - 0.80, 15) + 0.14)
            V.badge(slide, LEFT + 0.20, ry + 0.08, 0.32, "✓", fill=GREEN, size=11)
            V.text(slide, LEFT + 0.65, ry, lw - 0.85, rh, row, size=15, anchor=MSO_ANCHOR.MIDDLE)
            ry += rh + 0.06

        rx = LEFT + lw + 0.40
        rw = RIGHT - rx
        qh = max(1.05, _text_h(question, rw - 0.50, 16, True) + 0.62)
        V.box(slide, rx, cy, rw, qh, "", fill=LIGHT_NAVY, radius=0.08)
        V.text(slide, rx + 0.25, cy + 0.10, rw - 0.50, 0.32, "NEXT QUESTION", size=13,
               bold=True, color=RED)
        V.text(slide, rx + 0.25, cy + 0.42, rw - 0.50, qh - 0.50, question, size=16, bold=True,
               color=NAVY)
        ly = cy + qh + 0.22
        _heading(slide, rx, ly, rw, "In this part", NAVY)
        row = min(0.44, (y1 - ly - 0.40) / max(1, len(covers)))
        for k, cover in enumerate(covers):
            iy = ly + 0.40 + k * row
            V.badge(slide, rx, iy + (row - 0.30) / 2, 0.30, str(k + 1), fill=NAVY, size=11)
            V.text(slide, rx + 0.42, iy, rw - 0.42, row, cover, size=14,
                   anchor=MSO_ANCHOR.MIDDLE)
        _finish(slide, n, takeaway, notes)
    return step


CODE_CHAR = 0.0081  # Consolas width per point, inches


def _code_fits(text, size, w):
    return max(len(line) for line in text.split("\n")) * size * CODE_CHAR <= w - 0.38


def _code_items(blocks, size):
    return [{"code": b["code"], "label": b.get("label"), "kind": b.get("kind", "info"),
             "size": b.get("size", size)} for b in blocks]


def code(title, subtitle, blocks, points=(), takeaway=None, notes=None, mode="columns",
         left_w=7.80):
    """Code blocks with labels (good / bad / info) and short explanation points.

    The code font is the largest size (13 down to 9) that fits the width and height.
    """
    def step(prs, layout):
        slide, top, n = _new(prs, layout, title, subtitle)
        y0, y1 = top + 0.10, _bottom(takeaway)
        avail = y1 - y0
        if mode == "left":
            items = _code_items(blocks, 9)
            for size in (13, 12, 11, 10, 9):
                trial = _code_items(blocks, size)
                total = sum(item_height(i, left_w) for i in trial) + GAP * (len(trial) - 1)
                if all(_code_fits(i["code"], i["size"], left_w) for i in trial) and \
                        total <= avail:
                    items = trial
                    break
            stack_items(slide, LEFT, y0, left_w, avail, items)
            rx = LEFT + left_w + 0.30
            stack_items(slide, rx, y0, RIGHT - rx, avail, points)
        else:
            gap = 0.25
            bw = (WIDTH - gap * (len(blocks) - 1)) / len(blocks)
            ph = 0.0
            if points:
                pw = (WIDTH - 0.20 * (len(points) - 1)) / len(points)
                ph = max(item_height(p, pw) for p in points) + 0.20
            room = avail - (ph + 0.25 if points else 0)
            items = _code_items(blocks, 9)
            for size in (13, 12, 11, 10, 9):
                trial = _code_items(blocks, size)
                ch = max(item_height(i, bw) for i in trial)
                if all(_code_fits(i["code"], i["size"], bw) for i in trial) and ch <= room:
                    items = trial
                    break
            natural = max(item_height(i, bw) for i in items) + 0.10
            ch = min(natural + max(0.0, room - natural) * 0.5, room)
            for k, item in enumerate(items):
                draw_item(slide, item, LEFT + k * (bw + gap), y0, bw, ch)
            if points:
                py = y0 + ch + 0.25
                row_items(slide, LEFT, py, WIDTH, y1 - py, points)
        _finish(slide, n, takeaway, notes)
    return step


def exercise(title, subtitle, *, tasks, expected, scenario=None, scenario_code=None,
             kind="PRACTICE EXERCISE", minutes=10, worksheet=None,
             expected_title="Expected outcome", takeaway=None, notes=None):
    """Scenario, tasks and expected outcome."""
    def step(prs, layout):
        slide, top, n = _new(prs, layout, title, subtitle)
        y0, y1 = top + 0.10, _bottom(takeaway)
        official = "OFFICIAL" in kind
        tag_w = 0.25 + len(kind) * 0.115
        V.box(slide, LEFT, y0, tag_w, 0.42, kind, fill=RED if official else PURPLE, size=13)
        V.box(slide, LEFT + tag_w + 0.15, y0, 1.25, 0.42, f"{minutes} MIN", fill=LIGHT_GRAY,
              color=DARK_GRAY, size=13)
        if worksheet:
            ww = min(7.2, 0.45 + len(worksheet) * 0.085)
            V.code(slide, RIGHT - ww, y0, ww, 0.42, worksheet, size=11,
                   align=PP_ALIGN.CENTER)
        cy = y0 + 0.62
        w1, w2, gap = 4.20, 4.25, 0.25
        w3 = WIDTH - w1 - w2 - 2 * gap
        x2, x3 = LEFT + w1 + gap, LEFT + w1 + w2 + 2 * gap
        need1 = (_text_h(scenario, w1 - 0.30, 14) + 0.20) if scenario else 0
        if scenario_code:
            lines = scenario_code.split("\n")
            need1 += len(lines) * 11 * 1.30 / 72 + 0.34
        need2 = sum(max(0.44, _text_h(t, w2 - 0.75, 14) + 0.12) + 0.06 for t in tasks)
        need3 = sum(max(0.42, _text_h(e, w3 - 0.70, 13) + 0.12) + 0.04 for e in expected)
        ch = min(y1 - cy, max(2.60, max(need1, need2, need3) + 0.58 + 0.30))

        body = V.card(slide, LEFT, cy, w1, ch, "Scenario", head_fill=NAVY, title_size=14)
        sy = body
        if scenario:
            sh = _text_h(scenario, w1 - 0.30, 14) + 0.12
            V.text(slide, LEFT + 0.15, sy, w1 - 0.30, sh, scenario, size=14)
            sy += sh + 0.08
        if scenario_code:
            lines = scenario_code.split("\n")
            size = 11 if max(len(line) for line in lines) <= 40 else 10
            V.code(slide, LEFT + 0.12, sy, w1 - 0.24, min(cy + ch - sy - 0.12,
                   len(lines) * size * 1.30 / 72 + 0.34), scenario_code, size=size)

        body = V.card(slide, x2, cy, w2, ch, "Your tasks", head_fill=PURPLE, title_size=14)
        ty = body
        for k, task in enumerate(tasks):
            th = max(0.44, _text_h(task, w2 - 0.75, 14) + 0.12)
            V.badge(slide, x2 + 0.15, ty + 0.06, 0.32, str(k + 1), fill=PURPLE, size=11)
            V.text(slide, x2 + 0.58, ty, w2 - 0.73, th, task, size=14, anchor=MSO_ANCHOR.MIDDLE)
            ty += th + 0.06

        body = V.card(slide, x3, cy, w3, ch, expected_title, head_fill=GREEN, title_size=14)
        ey = body
        for row in expected:
            eh = max(0.42, _text_h(row, w3 - 0.70, 13) + 0.12)
            V.badge(slide, x3 + 0.15, ey + 0.06, 0.28, "✓", fill=GREEN, size=10)
            V.text(slide, x3 + 0.52, ey, w3 - 0.67, eh, row, size=13, anchor=MSO_ANCHOR.MIDDLE)
            ey += eh + 0.04
        _finish(slide, n, takeaway, notes)
    return step


def knowledge(title, subtitle, questions, takeaway=None, notes=None):
    """Up to ten numbered questions in two columns; answers belong in the notes."""
    def step(prs, layout):
        slide, top, n = _new(prs, layout, title, subtitle)
        y0, y1 = top + 0.10, _bottom(takeaway)
        per_col = math.ceil(len(questions) / 2)
        gap = 0.40
        cw = (WIDTH - gap) / 2
        rh = (y1 - y0 - 0.12 * (per_col - 1)) / per_col
        for k, q in enumerate(questions):
            col, row = divmod(k, per_col)
            x, y = LEFT + col * (cw + gap), y0 + row * (rh + 0.12)
            V.box(slide, x, y, cw, rh, "", fill=CARD_BG, line=CARD_LINE, radius=0.08)
            V.badge(slide, x + 0.16, y + (rh - 0.42) / 2, 0.42, str(k + 1),
                    fill=[NAVY, PURPLE, TEAL, GREEN, ORANGE, RED][k % 6], size=13)
            V.text(slide, x + 0.75, y, cw - 0.90, rh, q, size=15, anchor=MSO_ANCHOR.MIDDLE)
        _finish(slide, n, takeaway, notes)
    return step


def compare(title, subtitle, headers, rows, col_widths, items=(), takeaway=None, notes=None,
            row_h=0.56, size=14, first_col_bold=True):
    """Comparison table with an optional side panel."""
    def step(prs, layout):
        slide, top, n = _new(prs, layout, title, subtitle)
        y0, y1 = top + 0.10, _bottom(takeaway)
        V.table(slide, LEFT, y0, col_widths, headers, rows, row_h=row_h, size=size,
                first_col_bold=first_col_bold)
        if items:
            rx = LEFT + sum(col_widths) + 0.35
            stack_items(slide, rx, y0, RIGHT - rx, y1 - y0, items)
        _finish(slide, n, takeaway, notes)
    return step


def custom(draw, title, subtitle, takeaway=None, notes=None):
    """``draw(slide, y0, y1)`` fills the body area."""
    def step(prs, layout):
        slide, top, n = _new(prs, layout, title, subtitle)
        draw(slide, top + 0.10, _bottom(takeaway))
        _finish(slide, n, takeaway, notes)
    return step


def story(title, subtitle, *, gave_title, gave, now, path_title, path, takeaway=None,
          notes=None):
    """Module opener: what the previous module gave us, and the path through this one.

    ``path`` is a list of (step, part label, colour); ``now`` is a (label, text) callout.
    """
    def draw(slide, y0, y1):
        lw = 5.10
        body = V.card(slide, LEFT, y0, lw, y1 - y0, gave_title, head_fill=GREEN, title_size=14)
        for k, row in enumerate(gave):
            ry = body + 0.05 + k * 0.62
            V.badge(slide, LEFT + 0.20, ry + 0.10, 0.32, "✓", fill=GREEN, size=11)
            V.text(slide, LEFT + 0.65, ry, lw - 0.85, 0.55, row, size=15,
                   anchor=MSO_ANCHOR.MIDDLE)
        V.callout(slide, LEFT + 0.15, y1 - 1.25, lw - 0.30, 1.05, now[0], now[1], size=14)

        rx = LEFT + lw + 0.40
        rw = RIGHT - rx
        y = V.section(slide, rx, y0, rw, path_title, icon="🧭", fill=NAVY)
        rh = (y1 - y) / len(path)
        for k, (text, part, fill) in enumerate(path):
            ry = y + k * rh
            V.badge(slide, rx, ry + (rh - 0.38) / 2, 0.38, str(k + 1), fill=fill, size=12)
            V.text(slide, rx + 0.50, ry, rw - 2.10, rh, text, size=15, anchor=MSO_ANCHOR.MIDDLE)
            V.box(slide, RIGHT - 1.50, ry + (rh - 0.40) / 2, 1.50, 0.40, part,
                  fill=LIGHT.get(fill, LIGHT_NAVY), color=fill, size=12)
    return custom(draw, title, subtitle, takeaway=takeaway, notes=notes)


def wrapup(title, subtitle, *, can, next_title, questions, bring, takeaway=None, notes=None):
    """Module summary: what learners can now do, the next module's questions, what to bring."""
    def draw(slide, y0, y1):
        lw = 6.10
        body = V.card(slide, LEFT, y0, lw, y1 - y0, "You can now", head_fill=GREEN, title_size=14)
        row = min(0.66, (y1 - body - 0.15) / len(can))
        for k, text in enumerate(can):
            ry = body + 0.05 + k * row
            V.badge(slide, LEFT + 0.20, ry + (row - 0.32) / 2, 0.32, "✓", fill=GREEN, size=11)
            V.text(slide, LEFT + 0.65, ry, lw - 0.85, row, text, size=15,
                   anchor=MSO_ANCHOR.MIDDLE)
        rx = LEFT + lw + 0.40
        rw = RIGHT - rx
        body = V.card(slide, rx, y0, rw, 3.05, next_title, head_fill=NAVY, title_size=13)
        for k, q in enumerate(questions):
            qy = body + 0.02 + k * 0.72
            V.badge(slide, rx + 0.20, qy + 0.12, 0.38, str(k + 1), fill=NAVY, size=13)
            V.text(slide, rx + 0.72, qy, rw - 0.90, 0.62, q, size=15, anchor=MSO_ANCHOR.MIDDLE)
        V.callout(slide, rx, y0 + 3.25, rw, y1 - y0 - 3.25, bring[0], bring[1], size=15)
    return custom(draw, title, subtitle, takeaway=takeaway, notes=notes)


def myths(title, subtitle, pairs, anti, *, remember, takeaway=None, notes=None):
    """Misconception → reality rows beside an anti-pattern checklist.

    ``pairs`` is a list of (misconception, reality); ``remember`` is a (label, text) callout.
    """
    def draw(slide, y0, y1):
        lw = 7.60
        V.text(slide, LEFT, y0, 3.6, 0.34, "MISCONCEPTION", size=13, bold=True, color=RED)
        V.text(slide, LEFT + 3.95, y0, 3.6, 0.34, "REALITY", size=13, bold=True, color=GREEN)
        rh = (y1 - y0 - 0.45 - 0.12 * (len(pairs) - 1)) / len(pairs)
        for k, (myth, fact) in enumerate(pairs):
            ry = y0 + 0.45 + k * (rh + 0.12)
            V.box(slide, LEFT, ry, 3.55, rh, myth, fill=LIGHT_RED, color=RED, size=13,
                  radius=0.08)
            V.arrow(slide, LEFT + 3.58, ry + rh / 2, LEFT + 3.92, ry + rh / 2, color=MUTED)
            V.box(slide, LEFT + 3.95, ry, lw - 3.95, rh, fact, fill=LIGHT_GREEN, color=GREEN,
                  size=13, bold=False, radius=0.08)
        rx = LEFT + lw + 0.40
        stack_items(slide, rx, y0, RIGHT - rx, y1 - y0, [
            {"checks": anti, "h": "Anti-patterns to avoid", "c": RED, "mark": "✕", "size": 13},
            {"callout": remember},
        ])
    return custom(draw, title, subtitle, takeaway=takeaway, notes=notes)


def key_terms(title, subtitle, terms, takeaway=None, notes=None):
    """Up to eight key terms with plain definitions, in two columns."""
    def draw(slide, y0, y1):
        cols = 2
        rows = math.ceil(len(terms) / cols)
        gx, gy = 0.30, 0.16
        cw = (WIDTH - gx) / cols
        rh = (y1 - y0 - gy * (rows - 1)) / rows
        colors = [NAVY, PURPLE, TEAL, GREEN, ORANGE, RED]
        for k, (term, definition) in enumerate(terms):
            row, col = divmod(k, cols)
            x, y = LEFT + col * (cw + gx), y0 + row * (rh + gy)
            color = colors[k % len(colors)]
            V.box(slide, x, y, cw, rh, "", fill=CARD_BG, line=CARD_LINE, radius=0.06)
            V.badge(slide, x + 0.16, y + 0.12, 0.40, str(k + 1), fill=color, size=12)
            V.text(slide, x + 0.70, y + 0.08, cw - 0.85, 0.46, term, size=15, bold=True, color=color,
                   anchor=MSO_ANCHOR.MIDDLE)
            V.text(slide, x + 0.70, y + 0.52, cw - 0.85, rh - 0.58, definition, size=13)
    return custom(draw, title, subtitle, takeaway=takeaway, notes=notes)


def full_forms(title, subtitle, items, takeaway=None, notes=None):
    """Every abbreviation in the module with its full form, alphabetical down the columns."""
    def draw(slide, y0, y1):
        n = len(items)
        cols = 3 if n <= 30 else 4
        rows = math.ceil(n / cols)
        gx = 0.25
        cw = (WIDTH - gx * (cols - 1)) / cols
        rh = min(0.62, (y1 - y0) / rows)
        pill = 1.10 if cols == 3 else 0.95
        size = 12 if cols == 3 else 11
        colors = [NAVY, PURPLE, TEAL, GREEN]
        for k, (abbr, full) in enumerate(items):
            col, row = divmod(k, rows)
            x, y = LEFT + col * (cw + gx), y0 + row * rh
            color = colors[col % len(colors)]
            V.box(slide, x, y + 0.05, pill, rh - 0.10, abbr, fill=LIGHT.get(color, LIGHT_NAVY),
                  color=color, size=12 if len(abbr) <= 5 else 10, radius=0.10, margin=0.04)
            V.text(slide, x + pill + 0.12, y, cw - pill - 0.14, rh, full, size=size,
                   anchor=MSO_ANCHOR.MIDDLE)
    return custom(draw, title, subtitle, takeaway=takeaway, notes=notes)
