#!/usr/bin/env python3
"""Native-shape diagram helpers for the new-content MD287 decks.

Everything here draws editable PowerPoint shapes (boxes, arrows, chevrons,
cylinders, cards, tables) in the house style of the existing day decks. Page
chrome, title block, key-takeaway bar, fonts and colours come straight from
mdb_deck_kit, which is imported read-only -- nothing in it is changed.

All positions and sizes are in inches.
"""
from __future__ import annotations

import colorsys
import math
import sys
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import mdb_deck_kit as K  # noqa: E402

RED, NAVY, BLACK, WHITE = K.RED, K.NAVY, K.BLACK, K.WHITE
GREEN, PURPLE, ORANGE, TEAL = K.GREEN, K.PURPLE, K.ORANGE, K.TEAL
INK, MUTED, GRAY = K.INK, K.MUTED, K.GRAY
CARD_BG, CARD_LINE, TAKEAWAY_BG = K.CARD_BG, K.CARD_LINE, K.TAKEAWAY_BG

LIGHT_RED = RGBColor(0xFD, 0xEC, 0xEA)
LIGHT_GREEN = RGBColor(0xE8, 0xF5, 0xE9)
LIGHT_NAVY = RGBColor(0xE8, 0xEE, 0xF7)
LIGHT_TEAL = RGBColor(0xE0, 0xF2, 0xF1)
LIGHT_PURPLE = RGBColor(0xEF, 0xEA, 0xF7)
LIGHT_GRAY = RGBColor(0xEE, 0xEE, 0xEE)
MID_GRAY = RGBColor(0xBD, 0xBD, 0xBD)
DARK_GRAY = RGBColor(0x42, 0x42, 0x42)

BODY_BOTTOM = 6.38   # lowest y body content may reach (takeaway bar sits below)
LEFT = 0.50
RIGHT = 12.80
WIDTH = RIGHT - LEFT

# Nirmala UI averages ~0.0073in per character per point (0.095in at 13pt).
CHAR_W = 0.0073
WARNINGS: list[str] = []
_current_slide_label = ""


def I(v: float) -> int:
    return int(Inches(v))


# Diagram fills are softened so they are easier on the eyes: same hue, less
# saturation, nudged slightly toward mid-lightness so white labels stay legible.
# Only shape fills, outlines and arrows use this -- never font colours, and never
# the house chrome (rail, title, footer, takeaway bar) drawn by mdb_deck_kit.
SATURATION = 0.60
LIGHTNESS_EASE = 0.14
CHARCOAL = RGBColor(0x3A, 0x3A, 0x3A)


def soften(color):
    if color is None:
        return None
    r, g, b = (int(str(color)[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    if l > 0.88:
        return color  # pale tints are already soft, and must match the takeaway bar
    if s < 0.05:
        # Pure black boxes become charcoal; greys are left alone.
        return CHARCOAL if l < 0.10 else color
    s *= SATURATION
    l += (0.5 - l) * LIGHTNESS_EASE
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return RGBColor(round(r * 255), round(g * 255), round(b * 255))


def icon_badge(slide, x, y, d, icon, fill):
    K._icon_badge(slide, I(x), I(y), I(d), icon, soften(fill))


def set_slide_label(label: str) -> None:
    global _current_slide_label
    _current_slide_label = label


# ---------------------------------------------------------------------------
# Text
# ---------------------------------------------------------------------------
def _lines(text) -> list:
    if text is None:
        return []
    if isinstance(text, (list, tuple)):
        return list(text)
    return str(text).split("\n")


def _estimate_height(lines: list, width: float, size: float, bold: bool) -> float:
    per_char = CHAR_W * size * (1.07 if bold else 1.0)
    total = 0
    for line in lines:
        s = line[0] if isinstance(line, tuple) else line
        sz = (line[1].get("size", size) if isinstance(line, tuple) else size)
        chars = max(1, int(max(width, 0.2) / (CHAR_W * sz * (1.07 if bold else 1.0))))
        total += max(1, math.ceil(len(s) / chars)) * sz * 1.22 / 72
    return total


def _check_fit(lines, width, height, size, bold, what):
    need = _estimate_height(lines, width, size, bold)
    if need > height + 0.02:
        preview = (lines[0][0] if lines and isinstance(lines[0], tuple) else (lines[0] if lines else ""))
        WARNINGS.append(f"{_current_slide_label}: {what} may overflow "
                        f"(needs {need:.2f}in, has {height:.2f}in): {str(preview)[:50]!r}")


def fill_text(tf, text, *, size, color=INK, bold=False, italic=False, align=PP_ALIGN.LEFT,
              font=K.BODY_FONT, space_after=0):
    """Write lines into a text frame. A line may be (text, {size,bold,color,italic,font})."""
    first = True
    for line in _lines(text):
        opts = {}
        if isinstance(line, tuple):
            line, opts = line
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = opts.get("align", align)
        if space_after:
            p.space_after = Pt(space_after)
        K.add_run(
            p, line, size=opts.get("size", size), bold=opts.get("bold", bold),
            italic=opts.get("italic", italic), color=opts.get("color", color),
            font=opts.get("font", font),
        )


def text(slide, x, y, w, h, value, *, size=14, color=INK, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=K.BODY_FONT, space_after=0):
    box = K.add_textbox(slide, I(x), I(y), I(w), I(h), anchor=anchor)
    fill_text(box.text_frame, value, size=size, color=color, bold=bold, italic=italic,
              align=align, font=font, space_after=space_after)
    _check_fit(_lines(value), w, h, size, bold, "text")
    return box


def rich(slide, x, y, w, h, runs, *, size=14, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """One paragraph of mixed runs: [(text, {bold,color,italic,font})]."""
    box = K.add_textbox(slide, I(x), I(y), I(w), I(h), anchor=anchor)
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    for value, opts in runs:
        K.add_run(p, value, size=opts.get("size", size), bold=opts.get("bold", False),
                  italic=opts.get("italic", False), color=opts.get("color", INK),
                  font=opts.get("font", K.BODY_FONT))
    _check_fit(["".join(r[0] for r in runs)], w, h, size, False, "rich text")
    return box


# ---------------------------------------------------------------------------
# Shapes
# ---------------------------------------------------------------------------
def box(slide, x, y, w, h, value="", *, fill=NAVY, color=WHITE, size=14, bold=True,
        kind=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, line_w=1.0, dash=False,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=K.BODY_FONT,
        radius=0.14, italic=False, margin=0.08):
    shp = slide.shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = soften(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = soften(line)
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = I(margin)
    tf.margin_top = tf.margin_bottom = I(0.03)
    if value:
        fill_text(tf, value, size=size, color=color, bold=bold, italic=italic,
                  align=align, font=font)
        _check_fit(_lines(value), w - 2 * margin, h - 0.06, size, bold, "box")
    return shp


def badge(slide, x, y, d, value, *, fill=RED, color=WHITE, size=14):
    return box(slide, x, y, d, d, value, fill=fill, color=color, size=size,
               kind=MSO_SHAPE.OVAL, margin=0.0)


def cylinder(slide, x, y, w, h, value, *, fill=NAVY, color=WHITE, size=13):
    shp = box(slide, x, y, w, h, "", fill=fill, kind=MSO_SHAPE.CAN)
    try:
        shp.adjustments[0] = 0.18
    except (IndexError, ValueError):
        pass
    # Keep the label in the barrel, below the elliptical lid.
    text(slide, x, y + h * 0.22, w, h * 0.74, value, size=size, color=color, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return shp


def cloud(slide, x, y, w, h, value, *, fill=LIGHT_GRAY, color=DARK_GRAY, size=13):
    return box(slide, x, y, w, h, value, fill=fill, color=color, size=size,
               kind=MSO_SHAPE.CLOUD, line=MID_GRAY, margin=0.15)


def arrow(slide, x1, y1, x2, y2, *, color=MUTED, width=1.75, dash=False, head=True):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    conn.line.color.rgb = soften(color)
    conn.line.width = Pt(width)
    if dash:
        conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if head:
        ln = conn.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("h", "med")
    return conn


def cross(slide, x, y, d=0.36):
    return badge(slide, x, y, d, "✕", fill=RED, size=13)


def check(slide, x, y, d=0.36):
    return badge(slide, x, y, d, "✓", fill=GREEN, size=13)


def chevrons(slide, x, y, w, h, labels, fills, *, size=13, color=WHITE, gap=0.04):
    """A process band: pentagon first, chevrons after."""
    n = len(labels)
    notch = h * 0.30
    seg = (w + notch * (n - 1) - gap * (n - 1)) / n
    shapes = []
    for i, label in enumerate(labels):
        kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        sx = x + i * (seg - notch + gap)
        shp = box(slide, sx, y, seg, h, "", fill=fills[i % len(fills)], kind=kind)
        try:
            shp.adjustments[0] = 0.30
        except (IndexError, ValueError):
            pass
        inner_left = sx + (notch if i else 0.10)
        text(slide, inner_left, y, seg - notch - (notch if i else 0.10) + 0.05, h, label,
             size=size, color=color, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        shapes.append((sx, seg))
    return shapes


def vflow(slide, x, y, w, labels, *, box_h, gap, fills, size=14, colors=None,
          arrow_color=MUTED):
    """Boxes top to bottom with arrows between. Returns the bottom edge."""
    for i, label in enumerate(labels):
        by = y + i * (box_h + gap)
        box(slide, x, by, w, box_h, label, fill=fills[i % len(fills)],
            color=(colors[i % len(colors)] if colors else WHITE), size=size)
        if i < len(labels) - 1:
            arrow(slide, x + w / 2, by + box_h + 0.03, x + w / 2, by + box_h + gap - 0.03,
                  color=arrow_color)
    return y + len(labels) * box_h + (len(labels) - 1) * gap


def hflow(slide, x, y, labels, *, box_w, box_h, gap, fills, size=13, colors=None,
          arrow_color=MUTED):
    """Boxes left to right with arrows between. Returns the right edge."""
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        box(slide, bx, y, box_w, box_h, label, fill=fills[i % len(fills)],
            color=(colors[i % len(colors)] if colors else WHITE), size=size)
        if i < len(labels) - 1:
            arrow(slide, bx + box_w + 0.03, y + box_h / 2, bx + box_w + gap - 0.03, y + box_h / 2,
                  color=arrow_color)
    return x + len(labels) * box_w + (len(labels) - 1) * gap


def chips(slide, x, y, w, labels, *, cols, h=0.5, gap=0.12, fill=LIGHT_NAVY, color=NAVY,
          size=13, bold=True):
    rows = math.ceil(len(labels) / cols)
    cw = (w - gap * (cols - 1)) / cols
    for i, label in enumerate(labels):
        r, c = divmod(i, cols)
        box(slide, x + c * (cw + gap), y + r * (h + gap), cw, h, label, fill=fill,
            color=color, size=size, bold=bold)
    return y + rows * h + (rows - 1) * gap


def section(slide, x, y, w, label, *, icon="📌", fill=NAVY, size=14):
    """Small red heading with an icon badge, matching the day-deck panel headings."""
    d = 0.34
    icon_badge(slide, x, y, d, icon, fill)
    text(slide, x + d + 0.08, y - 0.01, w - d - 0.08, 0.38, label.upper(), size=size,
         color=RED, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    return y + 0.46


def pill(slide, x, y, w, label, *, fill=RED, size=14, h=0.40):
    box(slide, x, y, w, h, label.upper(), fill=fill, size=size)
    return y + h + 0.12


def panel(slide, x, y, w, heading, bullets, *, icon="📌", badge_color=NAVY, size=15):
    """Heading + bullets from the day-deck kit. Returns the bottom edge in inches."""
    scale = size / K.BODY_PT
    used = K.add_panel_bullets(
        slide, I(x), I(y), I(w), icon=icon, heading=heading.upper(),
        bullets=[("•", b) for b in bullets], badge_color=soften(badge_color), scale=scale,
    )
    return y + used / 914400


def bullets(slide, x, y, w, h, items, *, size=15, marker="•", marker_color=INK, space=4):
    tb = K.add_textbox(slide, I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        K.add_run(p, f"{marker}  ", size=size, bold=marker != "•", color=marker_color)
        K._add_bullet_text(p, item, size)
    _check_fit([f"{marker}  {t}" for t in items], w, h, size, False, "bullets")
    return tb


def card(slide, x, y, w, h, title, *, head_fill=NAVY, body_fill=WHITE, head_h=0.46,
         size=14, title_size=14):
    """Rounded card with a coloured header band. Returns the body's top y."""
    box(slide, x, y, w, h, "", fill=body_fill, line=CARD_LINE, line_w=1.0, radius=0.06)
    head = box(slide, x, y, w, head_h, "", fill=head_fill, kind=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
    try:
        head.adjustments[0] = 0.22
    except (IndexError, ValueError):
        pass
    text(slide, x + 0.15, y, w - 0.3, head_h, title.upper(), size=title_size, color=WHITE,
         bold=True, anchor=MSO_ANCHOR.MIDDLE)
    return y + head_h + 0.12


def callout(slide, x, y, w, h, label, body, *, fill=TAKEAWAY_BG, size=15, label_color=RED):
    box(slide, x, y, w, h, "", fill=fill, radius=0.08)
    return rich(slide, x + 0.2, y, w - 0.4, h,
                [(label.upper() + "   ", {"bold": True, "color": label_color}), (body, {})],
                size=size, anchor=MSO_ANCHOR.MIDDLE)


def code(slide, x, y, w, h, value, *, size=14, fill=CARD_BG, color=INK, align=PP_ALIGN.LEFT):
    return box(slide, x, y, w, h, value, fill=fill, color=color, size=size, bold=False,
               font=K.CODE_FONT, line=CARD_LINE, align=align, margin=0.18, radius=0.06)


def table(slide, x, y, col_widths, headers, rows, *, row_h=0.46, size=15, header_fill=NAVY,
          first_col_bold=False):
    n_rows = len(rows) + 1
    frame = slide.shapes.add_table(n_rows, len(headers), I(x), I(y),
                                   I(sum(col_widths)), I(row_h * n_rows))
    tbl = frame.table
    for c, cw in enumerate(col_widths):
        tbl.columns[c].width = I(cw)
    for r in range(n_rows):
        tbl.rows[r].height = I(row_h)
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = soften(header_fill)
        cell.text_frame.text = ""
        K.add_run(cell.text_frame.paragraphs[0], h, size=size, bold=True, color=WHITE)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else CARD_BG
            cell.text_frame.text = ""
            K.add_run(cell.text_frame.paragraphs[0], str(value), size=size,
                      bold=first_col_bold and c == 0, color=INK)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _check_fit([str(value)], col_widths[c] - 0.2, row_h - 0.08, size, False, "table cell")
    return y + row_h * n_rows


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------
def content_slide(prs, layout, *, number, title, subtitle):
    """Blank house-style slide with the title block. Returns (slide, body_top_in)."""
    slide = prs.slides.add_slide(layout)
    K._strip_placeholders(slide)
    set_slide_label(f"slide {number}")
    top = K.add_title_block(slide, title=title, subtitle=subtitle)
    return slide, top / 914400 + 0.04


def finish(slide, *, number, takeaway=None, notes=None):
    if takeaway:
        K.add_key_takeaway(slide, takeaway)
    K.add_rail_and_footer(slide, number)
    if notes:
        K.set_notes(slide, notes)
